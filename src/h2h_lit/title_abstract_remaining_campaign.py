"""Bounded staging coordinator for remaining-corpus screening and category coding."""

from __future__ import annotations

import csv
import fcntl
import hashlib
import json
import os
import signal
import threading
import time
from collections import deque
from concurrent.futures import FIRST_COMPLETED, Future, ThreadPoolExecutor, wait
from itertools import islice
from pathlib import Path
from typing import Any, Iterable, Mapping

from h2h_lit.inference import InferenceProvider
from h2h_lit.openai_provider import OpenAIResponsesProvider
from h2h_lit.title_abstract_fast_track_batch import (
    Pricing,
    RunSettings,
    _maximum_attempt_cost,
    _normalize_usage,
    _usage_cost,
    load_jsonl,
    sha256_file,
)
from h2h_lit.title_abstract_fast_track_overnight import (
    Budget,
    RuntimeWindow,
    _error_category,
    atomic_json,
    create_json,
    json_hash,
    stable_id,
    utc_now,
)
from h2h_lit.title_abstract_fast_track_repair import (
    EVIDENCED,
    NOT_EVIDENCED,
    evidence_units_for,
    inference_input_for,
)
from h2h_lit.title_abstract_precision_rescreen import (
    _result_csv_row,
    _screen_record,
    result_group,
)

RUNNER_VERSION = "1.0.0"
CODING_SCHEMA_VERSION = "coding-1.0.0"
CODING_PROMPT_VERSION = "1.0.0"
ANNOTATION_STATES = ("PRESENT", "ABSENT", "UNCERTAIN")
ASSISTANCE_MODES = ("Algorithmic", "Adaptive", "Conversational", "Immersive")
VISUALIZATION_MODALITIES = ("Desktop 2D", "Large Display", "VR", "AR/MR", "CAVE")
TASKS = (
    "Navigation and Multiscale Orientation",
    "Comparison and Differentiation",
    "Selection, Filtering, and Precision Interaction",
    "Sensemaking and Hypothesis Development",
    "Coordination and Collaborative Reasoning",
)
WORKFLOW_SUPPORT = ("SUPPORTED", "UNSUPPORTED", "UNKNOWN")
USAGE_KEYS = (
    "input_tokens",
    "cached_input_tokens",
    "cache_write_tokens",
    "output_tokens",
    "reasoning_tokens",
)


def _coding_item_schema(labels: tuple[str, ...]) -> dict[str, Any]:
    return {
        "type": "array",
        "minItems": len(labels),
        "maxItems": len(labels),
        "items": {
            "type": "object",
            "additionalProperties": False,
            "properties": {
                "label": {"type": "string", "enum": list(labels)},
                "state": {"type": "string", "enum": list(ANNOTATION_STATES)},
                "evidence_status": {
                    "type": "string",
                    "enum": [EVIDENCED, NOT_EVIDENCED],
                },
                "evidence_ids": {"type": "array", "items": {"type": "string"}},
                "rationale": {"type": "string", "minLength": 1},
            },
            "required": ["label", "state", "evidence_status", "evidence_ids", "rationale"],
        },
    }


def taxonomy_coding_response_schema() -> dict[str, Any]:
    workflow = {
        "type": "object",
        "additionalProperties": False,
        "properties": {
            "decision": {"type": "string", "enum": list(WORKFLOW_SUPPORT)},
            "evidence_status": {"type": "string", "enum": [EVIDENCED, NOT_EVIDENCED]},
            "evidence_ids": {"type": "array", "items": {"type": "string"}},
            "rationale": {"type": "string", "minLength": 1},
        },
        "required": ["decision", "evidence_status", "evidence_ids", "rationale"],
    }
    return {
        "type": "object",
        "additionalProperties": False,
        "properties": {
            "assistance_modes": _coding_item_schema(ASSISTANCE_MODES),
            "visualization_modalities": _coding_item_schema(VISUALIZATION_MODALITIES),
            "tasks": _coding_item_schema(TASKS),
            "workflow_support_review": workflow,
            "overall_rationale": {"type": "string", "minLength": 1},
        },
        "required": [
            "assistance_modes",
            "visualization_modalities",
            "tasks",
            "workflow_support_review",
            "overall_rationale",
        ],
    }


def _validate_evidence_fields(
    item: Mapping[str, Any], units: Mapping[str, Mapping[str, Any]], field: str
) -> dict[str, Any]:
    ids = item.get("evidence_ids")
    if not isinstance(ids, list) or any(not isinstance(value, str) for value in ids):
        raise ValueError(f"{field}.evidence_ids must be strings")
    if len(ids) != len(set(ids)):
        raise ValueError(f"{field}.evidence_ids contain duplicates")
    unknown = sorted(set(ids) - set(units))
    if unknown:
        raise ValueError(f"{field}.evidence_ids contain nonexistent IDs: {unknown}")
    status = item.get("evidence_status")
    if status not in {EVIDENCED, NOT_EVIDENCED}:
        raise ValueError(f"{field}.evidence_status is invalid")
    if status == EVIDENCED and not ids:
        raise ValueError(f"{field}.EVIDENCED requires evidence IDs")
    if status == NOT_EVIDENCED and ids:
        raise ValueError(f"{field}.NOT_EVIDENCED requires no evidence IDs")
    rationale = item.get("rationale")
    if not isinstance(rationale, str) or not rationale.strip():
        raise ValueError(f"{field}.rationale must be non-empty")
    return {
        **dict(item),
        "rationale": rationale.strip(),
        "evidence": [dict(units[value]) for value in ids],
    }


def validate_taxonomy_coding_payload(payload: Any, record: Mapping[str, Any]) -> dict[str, Any]:
    expected = {
        "assistance_modes",
        "visualization_modalities",
        "tasks",
        "workflow_support_review",
        "overall_rationale",
    }
    if not isinstance(payload, dict) or set(payload) != expected:
        raise ValueError("coding response has invalid top-level fields")
    units = {unit["id"]: unit for unit in evidence_units_for(record)}

    def dimension(name: str, labels: tuple[str, ...]) -> list[dict[str, Any]]:
        rows = payload[name]
        if not isinstance(rows, list) or len(rows) != len(labels):
            raise ValueError(f"{name} must contain every frozen label exactly once")
        observed = [row.get("label") if isinstance(row, dict) else None for row in rows]
        if len(set(observed)) != len(observed) or set(observed) != set(labels):
            raise ValueError(f"{name} label coverage or uniqueness is invalid")
        validated = []
        for row in rows:
            if set(row) != {"label", "state", "evidence_status", "evidence_ids", "rationale"}:
                raise ValueError(f"{name}.{row.get('label')} has invalid fields")
            if row["state"] not in ANNOTATION_STATES:
                raise ValueError(f"{name}.{row['label']}.state is invalid")
            checked = _validate_evidence_fields(row, units, f"{name}.{row['label']}")
            if row["state"] == "PRESENT" and row["evidence_status"] != EVIDENCED:
                raise ValueError(f"{name}.{row['label']}.PRESENT requires evidence")
            validated.append(checked)
        return sorted(validated, key=lambda row: labels.index(row["label"]))

    workflow = payload["workflow_support_review"]
    if not isinstance(workflow, dict) or set(workflow) != {
        "decision", "evidence_status", "evidence_ids", "rationale"
    }:
        raise ValueError("workflow_support_review has invalid fields")
    if workflow["decision"] not in WORKFLOW_SUPPORT:
        raise ValueError("workflow_support_review.decision is invalid")
    workflow_checked = _validate_evidence_fields(workflow, units, "workflow_support_review")
    if workflow["decision"] in {"SUPPORTED", "UNSUPPORTED"} and workflow["evidence_status"] != EVIDENCED:
        raise ValueError("supported or unsupported workflow review requires evidence")
    rationale = payload["overall_rationale"]
    if not isinstance(rationale, str) or not rationale.strip():
        raise ValueError("overall_rationale must be non-empty")
    return {
        "assistance_modes": dimension("assistance_modes", ASSISTANCE_MODES),
        "visualization_modalities": dimension("visualization_modalities", VISUALIZATION_MODALITIES),
        "tasks": dimension("tasks", TASKS),
        "workflow_support_review": workflow_checked,
        "overall_rationale": rationale.strip(),
        "screening_outcome_changed": False,
    }


def choose_concurrency(
    current: int,
    *,
    clean_completions: int,
    rate_limit_headers: Mapping[str, str],
    throttled: bool,
) -> int:
    """Conservative documented ramp; missing capacity evidence caps the run at eight."""
    if throttled:
        return max(4, current // 2)
    token_limit = int(rate_limit_headers.get("x-ratelimit-limit-tokens", "0") or 0)
    request_limit = int(rate_limit_headers.get("x-ratelimit-limit-requests", "0") or 0)
    target = {4: 8, 8: 16, 16: 32}.get(current, current)
    if clean_completions < 40 or target == current:
        return current
    if target > 8 and (token_limit < 500_000 or request_limit < 500):
        return current
    return target


def _code_record(
    record: dict[str, Any],
    records_dir: Path,
    prompt: str,
    prompt_hash: str,
    settings: RunSettings,
    provider: InferenceProvider,
    budget: Budget,
    runtime: RuntimeWindow,
) -> dict[str, Any]:
    record_dir = records_dir / record["canonical_id"].replace(":", "_")
    record_dir.mkdir(exist_ok=True)
    final_path = record_dir / "result.json"
    if final_path.exists():
        return json.loads(final_path.read_text(encoding="utf-8"))
    snapshot = inference_input_for(record)
    input_hash = json_hash(snapshot)
    parameters = {
        "max_output_tokens": settings.max_output_tokens,
        "reasoning_effort": settings.reasoning_effort,
        "response_schema_version": CODING_SCHEMA_VERSION,
        "service_tier": settings.service_tier,
        "store": False,
        "structured_output": "title_abstract_taxonomy_coding_v1_0_0",
        "verbosity": settings.verbosity,
    }
    attempts: list[dict[str, Any]] = []
    for attempt_number in range(1, settings.retry_limit + 2):
        reservation_path = record_dir / f"attempt-{attempt_number:03d}.reservation.json"
        attempt_path = record_dir / f"attempt-{attempt_number:03d}.json"
        if reservation_path.exists() and not attempt_path.exists():
            result = {
                "status": "AMBIGUOUS", "canonical_id": record["canonical_id"],
                "failure_reason": "OUTSTANDING_RESERVATION_NOT_RESUBMITTED", "attempts": attempts,
            }
            create_json(final_path, result)
            return result
        if not runtime.allow_new_request():
            result = {
                "status": "UNPROCESSED_RUNTIME_STOP" if not attempts else "FAILED",
                "canonical_id": record["canonical_id"], "failure_reason": "RUNTIME_DRAIN_STOP_BEFORE_REQUEST",
                "attempts": attempts,
            }
            create_json(final_path, result)
            return result
        maximum = _maximum_attempt_cost(prompt, snapshot, settings)
        if not budget.reserve(maximum):
            result = {
                "status": "UNPROCESSED_BUDGET_STOP" if not attempts else "FAILED",
                "canonical_id": record["canonical_id"], "failure_reason": "CUMULATIVE_BUDGET_STOP_BEFORE_REQUEST",
                "attempts": attempts,
            }
            create_json(final_path, result)
            return result
        runtime.mark_first_request()
        request_id = stable_id("remaining-campaign-coding-v1", record["canonical_id"], input_hash, prompt_hash, str(attempt_number))
        create_json(reservation_path, {
            "request_id": request_id, "attempt_number": attempt_number,
            "maximum_reserved_cost_usd": maximum, "input_hash": input_hash,
        })
        started = time.monotonic()
        raw = ""
        metadata: dict[str, Any] = {}
        validated = None
        error = None
        try:
            raw = provider.generate(
                model=settings.model, prompt=prompt, input_snapshot=snapshot,
                parameters=parameters, request_id=request_id, attempt_number=attempt_number,
            )
            validated = validate_taxonomy_coding_payload(json.loads(raw), record)
        except Exception as exc:  # noqa: BLE001 - persisted evidence
            error = f"{type(exc).__name__}: {exc}"
        try:
            if isinstance(provider, OpenAIResponsesProvider) or hasattr(provider, "metadata_for"):
                metadata = dict(provider.metadata_for(request_id, attempt_number))
        except Exception as exc:  # noqa: BLE001
            error = error or f"metadata error: {type(exc).__name__}: {exc}"
            validated = None
        try:
            usage = _normalize_usage(metadata.get("provider_usage"))
        except (TypeError, ValueError) as exc:
            usage = {key: 0 for key in USAGE_KEYS}
            error = error or f"invalid provider usage: {exc}"
            validated = None
        actual = _usage_cost(usage, settings.pricing) if any(usage.values()) else maximum
        budget.settle(maximum, actual)
        attempt = {
            "request_id": request_id, "attempt_number": attempt_number,
            "status": "VALID" if validated is not None else "INVALID",
            "latency_seconds": max(0.0, time.monotonic() - started), "usage": usage,
            "estimated_billed_cost_usd": actual, "raw_response": raw,
            "validation_error": error, "provider_metadata": metadata,
        }
        create_json(attempt_path, attempt)
        attempts.append(attempt)
        if validated is not None:
            result = {
                "status": "VALIDATED", "canonical_id": record["canonical_id"],
                "coding_order": record.get("coding_order"), "source_screening_order": record.get("batch_order"),
                "title": record.get("title", ""), "abstract_missing": not bool(str(record.get("abstract", "")).strip()),
                "input_hash": input_hash, "prompt_hash": prompt_hash,
                "output_schema_version": CODING_SCHEMA_VERSION, "coding": validated,
                "attempts": attempts,
            }
            create_json(final_path, result)
            return result
        if _error_category({"attempts": attempts}) in {"AUTHENTICATION", "BILLING"}:
            break
    result = {
        "status": "FAILED", "failure_reason": "RETRIES_EXHAUSTED",
        "canonical_id": record["canonical_id"], "title": record.get("title", ""), "attempts": attempts,
    }
    create_json(final_path, result)
    return result


def _append_jsonl(path: Path, value: Mapping[str, Any]) -> None:
    with path.open("a", encoding="utf-8", newline="\n") as handle:
        handle.write(json.dumps(value, sort_keys=True, ensure_ascii=True) + "\n")


def _attempt_usage(result: Mapping[str, Any]) -> tuple[float, dict[str, int], int, float, dict[str, str]]:
    cost = 0.0
    usage = {key: 0 for key in USAGE_KEYS}
    retries = 0
    latency = 0.0
    headers: dict[str, str] = {}
    for attempt in result.get("attempts", []):
        cost += float(attempt.get("estimated_billed_cost_usd", 0.0))
        retries += int(int(attempt.get("attempt_number", 1)) > 1)
        latency += float(attempt.get("latency_seconds", 0.0))
        for key in USAGE_KEYS:
            usage[key] += int(attempt.get("usage", {}).get(key, 0))
        headers.update(attempt.get("provider_metadata", {}).get("rate_limit_headers", {}))
    return cost, usage, retries, latency, headers


class CampaignProgress:
    def __init__(self, output_dir: Path, queue_count: int, seed_count: int, initial_cost: float, initial_usage: Mapping[str, int]):
        self.output_dir = output_dir
        self.queue_count = queue_count
        self.seed_count = seed_count
        self.initial_cost = initial_cost
        self.initial_usage = {key: int(initial_usage.get(key, 0)) for key in USAGE_KEYS}
        self.screen = {name: 0 for name in ("advance", "deferred", "excluded", "background", "failed", "ambiguous", "unprocessed")}
        self.code = {name: 0 for name in ("validated", "failed", "ambiguous", "unprocessed")}
        self.stage_cost = {"screening": 0.0, "coding": 0.0}
        self.stage_usage = {stage: {key: 0 for key in USAGE_KEYS} for stage in ("screening", "coding")}
        self.requests = {"screening": 0, "coding": 0}
        self.retries = {"screening": 0, "coding": 0}
        self.latency = {"screening": 0.0, "coding": 0.0}
        self.completed_screen_ids: set[str] = set()
        self.completed_code_ids: set[str] = set()
        self.e7_warnings = 0
        self.rate_limit_headers: dict[str, str] = {}
        self.lock = threading.Lock()

    def add(self, stage: str, result: Mapping[str, Any]) -> None:
        record_id = str(result["canonical_id"])
        with self.lock:
            completed = self.completed_screen_ids if stage == "screening" else self.completed_code_ids
            if record_id in completed:
                return
            completed.add(record_id)
            if stage == "screening":
                group = result_group(result)
                self.screen[group] += 1
                if result.get("judgment", {}).get("route") == "H2H3_BACKGROUND":
                    self.screen["background"] += 1
                self.e7_warnings += int(bool(result.get("e7_consistency_warning")))
            else:
                status = result.get("status")
                self.code[
                    "validated" if status == "VALIDATED"
                    else "ambiguous" if status == "AMBIGUOUS"
                    else "unprocessed" if str(status).startswith("UNPROCESSED")
                    else "failed"
                ] += 1
            cost, usage, retries, latency, headers = _attempt_usage(result)
            self.stage_cost[stage] += cost
            self.requests[stage] += len(result.get("attempts", []))
            self.retries[stage] += retries
            self.latency[stage] += latency
            for key in USAGE_KEYS:
                self.stage_usage[stage][key] += usage[key]
            self.rate_limit_headers.update(headers)

    def snapshot(self, budget: Budget, runtime: RuntimeWindow, stop_reason: str | None, concurrency: int) -> dict[str, Any]:
        with self.lock:
            state = budget.snapshot()
            screen_done = len(self.completed_screen_ids)
            code_done = len(self.completed_code_ids)
            return {
                "updated_at_utc": utc_now(), "status": "RUNNING" if stop_reason is None else "STOPPING",
                "stop_reason": stop_reason, "current_concurrency": concurrency,
                "screening": {
                    "queue_records": self.queue_count, "completed_records": screen_done,
                    "unprocessed_records": self.queue_count - screen_done, "counts": dict(self.screen),
                    "usage": dict(self.stage_usage["screening"]), "estimated_cost_usd": self.stage_cost["screening"],
                    "requests": self.requests["screening"], "retries": self.retries["screening"],
                    "aggregate_latency_seconds": self.latency["screening"], "e7_consistency_warnings": self.e7_warnings,
                },
                "coding": {
                    "initial_seed_records": self.seed_count, "completed_records": code_done,
                    "counts": dict(self.code), "usage": dict(self.stage_usage["coding"]),
                    "estimated_cost_usd": self.stage_cost["coding"], "requests": self.requests["coding"],
                    "retries": self.retries["coding"], "aggregate_latency_seconds": self.latency["coding"],
                },
                "historical_cumulative_conservative_cost_usd": self.initial_cost,
                "cumulative_conservative_cost_usd": state["actual"], "outstanding_reservations_usd": state["reserved"],
                "hard_cumulative_spending_cap_usd": state["cap"], "elapsed_seconds": runtime.elapsed(),
                "combined_completions_per_hour": ((screen_done + code_done) / runtime.elapsed() * 3600 if runtime.elapsed() else 0.0),
                "latest_rate_limit_headers": dict(self.rate_limit_headers),
            }


def _coding_csv_row(result: Mapping[str, Any], record: Mapping[str, Any]) -> dict[str, Any]:
    coding = result.get("coding", {})
    def labels(name: str, state: str) -> str:
        return "; ".join(row["label"] for row in coding.get(name, []) if row.get("state") == state)
    evidence = []
    for name in ("assistance_modes", "visualization_modalities", "tasks"):
        for row in coding.get(name, []):
            if row.get("state") == "PRESENT":
                evidence.append(f"{name}:{row['label']} <- {','.join(row.get('evidence_ids', []))}")
    workflow = coding.get("workflow_support_review", {})
    return {
        "canonical_id": record["canonical_id"], "title": record.get("title", ""),
        "abstract": record.get("abstract", ""), "doi": record.get("doi") or "",
        "source_url": record.get("source_url") or "", "abstract_status": "MISSING" if not str(record.get("abstract", "")).strip() else "PRESENT",
        "assistance_present": labels("assistance_modes", "PRESENT"), "assistance_unknown": labels("assistance_modes", "UNCERTAIN"),
        "modalities_present": labels("visualization_modalities", "PRESENT"), "modalities_unknown": labels("visualization_modalities", "UNCERTAIN"),
        "tasks_present": labels("tasks", "PRESENT"), "tasks_unknown": labels("tasks", "UNCERTAIN"),
        "workflow_support_review": workflow.get("decision", ""), "category_evidence_ids": " | ".join(evidence),
        "coding_rationale": coding.get("overall_rationale", ""), "provisional_title_abstract_coding": "YES",
        "human_category_review": "", "human_final_eligibility": "", "human_notes": "",
    }


CODING_CSV_FIELDS = list(_coding_csv_row({}, {"canonical_id": ""}))


def _write_workbook(path: Path, csv_path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    workbook = Workbook()
    info = workbook.active
    info.title = "Instructions"
    info.append(["Provisional abstract-based category coding; not final eligibility or full-report extraction."])
    info.append(["Category UNKNOWN/UNSUPPORTED flags do not rewrite preserved screening outcomes."])
    sheet = workbook.create_sheet("Coded candidates")
    with csv_path.open(encoding="utf-8", newline="") as handle:
        for row in csv.reader(handle):
            sheet.append(row)
    fill = PatternFill("solid", fgColor="263B50")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for field, formula in (
        ("human_category_review", '"CONFIRMED,REVISED,UNCERTAIN"'),
        ("human_final_eligibility", '"INCLUDE,EXCLUDE,DEFER,NOT_ASSESSED"'),
    ):
        column = CODING_CSV_FIELDS.index(field) + 1
        validation = DataValidation(type="list", formula1=formula, allow_blank=True)
        sheet.add_data_validation(validation)
        validation.add(f"{sheet.cell(2, column).coordinate}:{sheet.cell(max(2, sheet.max_row), column).coordinate}")
    for column in range(1, sheet.max_column + 1):
        header = str(sheet.cell(1, column).value)
        sheet.column_dimensions[sheet.cell(1, column).column_letter].width = 80 if header in {"abstract", "coding_rationale"} else 28
    workbook.save(path)


def _iter_after(path: Path, skip: int) -> Iterable[dict[str, Any]]:
    for index, row in enumerate(load_jsonl(path)):
        if index >= skip:
            yield row


def run_remaining_campaign(
    *,
    queue_path: Path,
    coding_seed_path: Path,
    manifest_path: Path,
    config_path: Path,
    screening_prompt_path: Path,
    coding_prompt_path: Path,
    protocol_path: Path,
    amendment_path: Path,
    campaign_ledger_path: Path,
    output_dir: Path,
    settings: RunSettings,
    provider: InferenceProvider,
    runtime_hours: float = 12.0,
    drain_seconds: float = 300.0,
    resume: bool = False,
) -> dict[str, Any]:
    settings.validate()
    if settings.model != "gpt-5.6-luna" or settings.service_tier != "default":
        raise ValueError("authorized model/service-tier binding changed")
    if settings.smoke_count != 20 or settings.retry_limit > 1 or settings.concurrency != 4:
        raise ValueError("campaign requires 20-record smoke, concurrency 4, and at most one retry")
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    config = json.loads(config_path.read_text(encoding="utf-8"))
    queue_count = int(manifest.get("screening_queue", {}).get("records", 0))
    coding_seed_count = int(manifest.get("coding_seed", {}).get("records", 0))
    if queue_count < settings.smoke_count:
        raise ValueError("screening queue is smaller than the required smoke gate")
    if coding_seed_count < 1:
        raise ValueError("coding seed must contain at least one pending candidate")
    if manifest.get("screening_queue", {}).get("sha256") != sha256_file(queue_path):
        raise ValueError("remaining screening queue binding drift")
    if manifest.get("coding_seed", {}).get("sha256") != sha256_file(coding_seed_path):
        raise ValueError("coding seed binding drift")
    for name, path in (
        ("config", config_path), ("screening_prompt", screening_prompt_path),
        ("coding_prompt", coding_prompt_path), ("protocol", protocol_path), ("amendment", amendment_path),
    ):
        if manifest.get("bindings", {}).get(name, {}).get("sha256") != sha256_file(path):
            raise ValueError(f"{name} binding drift")
    if config.get("taxonomy") != {
        "annotation_states": list(ANNOTATION_STATES), "assistance_modes": list(ASSISTANCE_MODES),
        "visualization_modalities": list(VISUALIZATION_MODALITIES), "tasks": list(TASKS),
        "workflow_support_review": list(WORKFLOW_SUPPORT),
    }:
        raise ValueError("frozen taxonomy drift")
    screening_prompt = screening_prompt_path.read_text(encoding="utf-8")
    coding_prompt = coding_prompt_path.read_text(encoding="utf-8")
    if "Prompt-Version: 2.2.0" not in screening_prompt or f"Prompt-Version: {CODING_PROMPT_VERSION}" not in coding_prompt:
        raise ValueError("prompt version drift")
    ledger = json.loads(campaign_ledger_path.read_text(encoding="utf-8"))
    initial_actual = float(ledger["cumulative_conservative_cost_usd"])
    carried_reservation = max(0.0, float(ledger.get("current_inflight_reserved_cost_usd", ledger.get("outstanding_reservations_usd", 0.0))))
    initial_cost = initial_actual + carried_reservation
    initial_usage = ledger.get("cumulative_usage", {})
    if initial_cost > settings.hard_spending_cap_usd:
        raise ValueError("persisted campaign cost exceeds authorized cap")
    if output_dir.exists() and not resume:
        raise FileExistsError("output directory exists; explicit resume required")
    output_dir.mkdir(parents=True, exist_ok=True)
    screen_dir = output_dir / "screening_records"
    code_dir = output_dir / "coding_records"
    screen_dir.mkdir(exist_ok=True)
    code_dir.mkdir(exist_ok=True)
    lock_handle = (output_dir / "budget_owner.lock").open("a+", encoding="utf-8")
    try:
        fcntl.flock(lock_handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
    except BlockingIOError as exc:
        raise RuntimeError("another worker owns the shared campaign budget") from exc

    binding = {
        "artifact_class": "remaining_corpus_screening_and_coding_run", "runner_version": RUNNER_VERSION,
        "status": "RUNNING", "created_at_utc": utc_now(),
        "queue": {"path": str(queue_path), "sha256": sha256_file(queue_path), "records": queue_count},
        "coding_seed": {"path": str(coding_seed_path), "sha256": sha256_file(coding_seed_path), "records": coding_seed_count},
        "manifest": {"path": str(manifest_path), "sha256": sha256_file(manifest_path)},
        "config": {"path": str(config_path), "sha256": sha256_file(config_path)},
        "prompts": {"screening": sha256_file(screening_prompt_path), "coding": sha256_file(coding_prompt_path)},
        "protocol": {"path": str(protocol_path), "sha256": sha256_file(protocol_path), "version": "2.0.0"},
        "amendment": {"path": str(amendment_path), "sha256": sha256_file(amendment_path), "version": "2.2.0"},
        "schemas": {"screening": "2.2.0", "coding": CODING_SCHEMA_VERSION, "coding_sha256": json_hash(taxonomy_coding_response_schema())},
        "implementation": {"path": str(Path(__file__)), "sha256": sha256_file(Path(__file__))},
        "campaign_ledger": {"path": str(campaign_ledger_path), "sha256": sha256_file(campaign_ledger_path),
            "starting_actual_usd": initial_actual, "carried_unresolved_reservations_usd": carried_reservation},
        "provider": settings.provider_name, "model": settings.model,
        "parameters": {"reasoning_effort": settings.reasoning_effort, "service_tier": settings.service_tier,
            "max_output_tokens": settings.max_output_tokens, "verbosity": settings.verbosity, "store": False, "tools": []},
        "pricing_usd_per_million_tokens": {"input": settings.pricing.input_usd_per_million,
            "cached_input": settings.pricing.cached_input_usd_per_million,
            "cache_write": settings.pricing.cache_write_usd_per_million, "output": settings.pricing.output_usd_per_million},
        "controls": {"hard_cumulative_cap_usd": settings.hard_spending_cap_usd, "runtime_hours": runtime_hours,
            "drain_seconds": drain_seconds, "concurrency_ramp": [4, 8, 16, 32], "retry_limit": settings.retry_limit,
            "screening_smoke": 20, "coding_smoke": 5, "scheduling": "3 screening : 1 coding when available"},
        "prior_model_rationales_used_as_coding_evidence": False, "staging_only": True,
    }
    binding_path = output_dir / "run_binding.json"
    if binding_path.exists():
        previous = json.loads(binding_path.read_text(encoding="utf-8"))
        for key in ("queue", "coding_seed", "manifest", "config", "prompts", "schemas", "model", "parameters", "pricing_usd_per_million_tokens", "controls"):
            if previous.get(key) != binding.get(key):
                raise ValueError(f"resume binding mismatch: {key}")
        binding = previous
    else:
        create_json(binding_path, binding)

    def existing_results(directory: Path) -> dict[str, dict[str, Any]]:
        found = {}
        for path in directory.glob("*/result.json"):
            value = json.loads(path.read_text(encoding="utf-8"))
            record_id = value.get("canonical_id")
            if not isinstance(record_id, str) or record_id in found:
                raise ValueError("duplicate or malformed persisted result")
            found[record_id] = value
        return found

    existing_screen = existing_results(screen_dir)
    existing_code = existing_results(code_dir)
    persisted_cost = sum(_attempt_usage(row)[0] for row in (*existing_screen.values(), *existing_code.values()))
    budget = Budget(settings.hard_spending_cap_usd, initial_cost + persisted_cost)
    runtime = RuntimeWindow(output_dir, runtime_hours, drain_seconds)
    progress = CampaignProgress(output_dir, queue_count, coding_seed_count, initial_cost, initial_usage)
    for row in existing_screen.values():
        progress.add("screening", row)
    for row in existing_code.values():
        progress.add("coding", row)

    screen_csv = {name: output_dir / f"screening_{name}.csv" for name in ("advance", "deferred", "excluded", "background", "failed", "ambiguous", "unprocessed")}
    screen_fields = list(_result_csv_row({}, {}))
    for path in screen_csv.values():
        if not path.exists():
            with path.open("x", encoding="utf-8", newline="") as handle:
                csv.DictWriter(handle, fieldnames=screen_fields).writeheader()
    coding_csv = output_dir / "categorized_candidates.csv"
    if not coding_csv.exists():
        with coding_csv.open("x", encoding="utf-8", newline="") as handle:
            csv.DictWriter(handle, fieldnames=CODING_CSV_FIELDS).writeheader()
    unsupported_csv = output_dir / "coding_workflow_review_flags.csv"
    if not unsupported_csv.exists():
        with unsupported_csv.open("x", encoding="utf-8", newline="") as handle:
            csv.DictWriter(handle, fieldnames=CODING_CSV_FIELDS).writeheader()
    code_source: dict[str, dict[str, Any]] = {row["canonical_id"]: row for row in load_jsonl(coding_seed_path)}
    for result in existing_code.values():
        if result["canonical_id"] not in code_source:
            source_path = output_dir / "new_advance_coding_queue.jsonl"
            if source_path.exists():
                for row in load_jsonl(source_path):
                    code_source.setdefault(row["canonical_id"], row)

    stop_event = threading.Event()
    stop_reason: str | None = None
    stage_paused = {"screening": False, "coding": False}
    rolling = {stage: deque(maxlen=20) for stage in stage_paused}
    persistent_errors = deque(maxlen=3)
    current_concurrency = 4
    clean_since_ramp = 0

    def request_stop(reason: str) -> None:
        nonlocal stop_reason
        if stop_reason is None:
            stop_reason = reason
            stop_event.set()
            print(json.dumps({"event": "stop_requested", "reason": reason, "at": utc_now()}), flush=True)

    def signal_stop(signum: int, _frame: Any) -> None:
        request_stop(f"SIGNAL_{signum}_GRACEFUL_DRAIN")

    old_handlers = {sig: signal.signal(sig, signal_stop) for sig in (signal.SIGTERM, signal.SIGINT)}

    def persist(stage: str, result: dict[str, Any], source: Mapping[str, Any]) -> None:
        nonlocal current_concurrency, clean_since_ramp
        progress.add(stage, result)
        _append_jsonl(output_dir / f"{stage}_completion_events.jsonl", {
            "completed_at_utc": utc_now(), "stage": stage, "canonical_id": result.get("canonical_id"),
            "status": result.get("status"),
        })
        exhausted = result.get("status") == "FAILED" and result.get("failure_reason") == "RETRIES_EXHAUSTED"
        rolling[stage].append(exhausted)
        if len(rolling[stage]) == 20 and sum(rolling[stage]) >= 3:
            stage_paused[stage] = True
            print(json.dumps({"event": "stage_paused", "stage": stage, "reason": "ROLLING_VALIDATION_FAILURE_GATE"}), flush=True)
            if all(stage_paused.values()):
                request_stop("BOTH_STAGES_PAUSED")
        category = _error_category(result)
        if category in {"AUTHENTICATION", "BILLING"}:
            request_stop(f"{category}_FAILURE")
        if category in {"RATE_LIMIT", "STORAGE"}:
            persistent_errors.append(category)
            current_concurrency = choose_concurrency(current_concurrency, clean_completions=0, rate_limit_headers={}, throttled=True)
            if len(persistent_errors) == 3 and len(set(persistent_errors)) == 1:
                request_stop(f"PERSISTENT_{category}_FAILURE")
        else:
            persistent_errors.clear()
            if result.get("status") == "VALIDATED":
                clean_since_ramp += 1
                advanced = choose_concurrency(current_concurrency, clean_completions=clean_since_ramp,
                    rate_limit_headers=progress.rate_limit_headers, throttled=False)
                if advanced != current_concurrency:
                    current_concurrency = advanced
                    clean_since_ramp = 0
                    print(json.dumps({"event": "concurrency_ramped", "concurrency": current_concurrency}), flush=True)
        if stage == "screening":
            group = result_group(result)
            if group in screen_csv:
                with screen_csv[group].open("a", encoding="utf-8", newline="") as handle:
                    csv.DictWriter(handle, fieldnames=screen_fields).writerow(_result_csv_row(result, source))
            if result.get("judgment", {}).get("route") == "H2H3_BACKGROUND":
                with screen_csv["background"].open("a", encoding="utf-8", newline="") as handle:
                    csv.DictWriter(handle, fieldnames=screen_fields).writerow(_result_csv_row(result, source))
        elif result.get("status") == "VALIDATED":
            row = _coding_csv_row(result, source)
            with coding_csv.open("a", encoding="utf-8", newline="") as handle:
                csv.DictWriter(handle, fieldnames=CODING_CSV_FIELDS).writerow(row)
            if result.get("coding", {}).get("workflow_support_review", {}).get("decision") != "SUPPORTED":
                with unsupported_csv.open("a", encoding="utf-8", newline="") as handle:
                    csv.DictWriter(handle, fieldnames=CODING_CSV_FIELDS).writerow(row)
        atomic_json(output_dir / "progress.json", progress.snapshot(budget, runtime, stop_reason, current_concurrency))

    screen_prompt_hash = hashlib.sha256(screening_prompt.encode()).hexdigest()
    code_prompt_hash = hashlib.sha256(coding_prompt.encode()).hexdigest()

    def execute_smoke(stage: str, rows: list[dict[str, Any]]) -> bool:
        fn = _screen_record if stage == "screening" else _code_record
        directory = screen_dir if stage == "screening" else code_dir
        prompt = screening_prompt if stage == "screening" else coding_prompt
        prompt_hash = screen_prompt_hash if stage == "screening" else code_prompt_hash
        with ThreadPoolExecutor(max_workers=4) as pool:
            futures = {pool.submit(fn, row, directory, prompt, prompt_hash, settings, provider, budget, runtime): row for row in rows}
            for future in futures:
                result = future.result()
                persist(stage, result, futures[future])
                if result.get("status") != "VALIDATED":
                    request_stop(f"{stage.upper()}_SMOKE_VALIDATION_FAILURE")
        return not stop_event.is_set() and len(rows) > 0

    smoke_screen = list(islice(_iter_after(queue_path, 0), 20))
    smoke_code = list(islice(_iter_after(coding_seed_path, 0), 5))
    screening_smoke_passed = False
    coding_smoke_passed = False
    try:
        screening_smoke_passed = execute_smoke("screening", smoke_screen)
        if screening_smoke_passed:
            coding_smoke_passed = execute_smoke("coding", smoke_code)
        if screening_smoke_passed and coding_smoke_passed and not stop_event.is_set():
            print(json.dumps({
                "event": "smoke_gate_passed",
                "screening": len(smoke_screen),
                "coding": len(smoke_code),
            }), flush=True)
            current_concurrency = 8
            screening_iter = iter(_iter_after(queue_path, 20))
            coding_pending = deque(row for row in _iter_after(coding_seed_path, 5) if row["canonical_id"] not in existing_code)
            screen_exhausted = False
            coding_sequence = max(
                (int(row.get("coding_order") or 0) for row in code_source.values()),
                default=0,
            )
            for source in smoke_screen:
                result_path = (
                    screen_dir
                    / source["canonical_id"].replace(":", "_")
                    / "result.json"
                )
                result = json.loads(result_path.read_text(encoding="utf-8"))
                if (
                    result.get("status") == "VALIDATED"
                    and result.get("judgment", {}).get("operational_disposition")
                    == "ADVANCE_TO_FULL_REPORT_ASSESSMENT"
                    and source["canonical_id"] not in code_source
                ):
                    coding_sequence += 1
                    candidate = {
                        **source,
                        "coding_order": coding_sequence,
                        "coding_source": "new_remaining_screening_advance",
                    }
                    code_source[candidate["canonical_id"]] = candidate
                    _append_jsonl(
                        output_dir / "new_advance_coding_queue.jsonl", candidate
                    )
                    coding_pending.append(candidate)
            schedule_counter = 0
            futures: dict[Future[dict[str, Any]], tuple[str, dict[str, Any]]] = {}
            with ThreadPoolExecutor(max_workers=32) as pool:
                while True:
                    if not runtime.allow_new_request() and not stop_event.is_set():
                        request_stop("RUNTIME_LIMIT_DRAIN")
                    while not stop_event.is_set() and len(futures) < current_concurrency:
                        stage = "coding" if coding_pending and schedule_counter % 4 == 3 and not stage_paused["coding"] else "screening"
                        row = None
                        if stage == "screening" and not stage_paused["screening"] and not screen_exhausted:
                            try:
                                row = next(screening_iter)
                            except StopIteration:
                                screen_exhausted = True
                        if row is None and coding_pending and not stage_paused["coding"]:
                            stage = "coding"
                            row = coding_pending.popleft()
                        if row is None and not screen_exhausted and not stage_paused["screening"]:
                            stage = "screening"
                            try:
                                row = next(screening_iter)
                            except StopIteration:
                                screen_exhausted = True
                        if row is None:
                            break
                        if row["canonical_id"] in (existing_screen if stage == "screening" else existing_code):
                            continue
                        if stage == "screening":
                            future = pool.submit(_screen_record, row, screen_dir, screening_prompt, screen_prompt_hash, settings, provider, budget, runtime)
                        else:
                            code_source[row["canonical_id"]] = row
                            future = pool.submit(_code_record, row, code_dir, coding_prompt, code_prompt_hash, settings, provider, budget, runtime)
                        futures[future] = (stage, row)
                        schedule_counter += 1
                    if not futures:
                        break
                    done, _ = wait(futures, return_when=FIRST_COMPLETED)
                    for future in done:
                        stage, source = futures.pop(future)
                        result = future.result()
                        persist(stage, result, source)
                        if stage == "screening" and result.get("status") == "VALIDATED" and result.get("judgment", {}).get("operational_disposition") == "ADVANCE_TO_FULL_REPORT_ASSESSMENT":
                            coding_sequence += 1
                            candidate = {**source, "coding_order": coding_sequence, "coding_source": "new_remaining_screening_advance"}
                            code_source[candidate["canonical_id"]] = candidate
                            _append_jsonl(output_dir / "new_advance_coding_queue.jsonl", candidate)
                            coding_pending.append(candidate)
                        print(json.dumps({"event": "record_completed", "stage": stage, "canonical_id": result.get("canonical_id"), "status": result.get("status")}), flush=True)
    finally:
        for sig, handler in old_handlers.items():
            signal.signal(sig, handler)

    existing_screen = existing_results(screen_dir)
    existing_code = existing_results(code_dir)
    unprocessed_path = output_dir / "screening_unprocessed.jsonl"
    with unprocessed_path.open("w", encoding="utf-8", newline="\n") as handle:
        for row in load_jsonl(queue_path):
            result = existing_screen.get(row["canonical_id"])
            if result is None or str(result.get("status", "")).startswith("UNPROCESSED"):
                handle.write(json.dumps(row, sort_keys=True) + "\n")
    uncoded_path = output_dir / "uncoded_survivors.jsonl"
    with uncoded_path.open("w", encoding="utf-8", newline="\n") as handle:
        for row in code_source.values():
            result = existing_code.get(row["canonical_id"])
            if result is None or str(result.get("status", "")).startswith("UNPROCESSED"):
                handle.write(json.dumps(row, sort_keys=True) + "\n")

    category_summary = {"unique_coded_papers": 0, "multilabel_present_assignments": 0, "labels": {}, "workflow_support_review": {value: 0 for value in WORKFLOW_SUPPORT}}
    for result in existing_code.values():
        if result.get("status") != "VALIDATED":
            continue
        category_summary["unique_coded_papers"] += 1
        coding = result["coding"]
        for name in ("assistance_modes", "visualization_modalities", "tasks"):
            for row in coding[name]:
                key = f"{name}:{row['label']}:{row['state']}"
                category_summary["labels"][key] = category_summary["labels"].get(key, 0) + 1
                if row["state"] == "PRESENT":
                    category_summary["multilabel_present_assignments"] += 1
        category_summary["workflow_support_review"][coding["workflow_support_review"]["decision"]] += 1
    atomic_json(output_dir / "category_summary.json", category_summary)
    workbook_path = output_dir / "categorized_candidates.xlsx"
    _write_workbook(workbook_path, coding_csv)
    snapshot = progress.snapshot(budget, runtime, stop_reason, current_concurrency)
    snapshot.update({
        "status": "COMPLETE" if len(existing_screen) == queue_count and not any(
            row["canonical_id"] not in existing_code
            or str(existing_code[row["canonical_id"]].get("status", "")).startswith("UNPROCESSED")
            for row in code_source.values()
        ) and not stop_reason else "STOPPED",
        "stopping_reason": stop_reason or "AVAILABLE_WORK_COMPLETE", "screening_smoke_passed": screening_smoke_passed,
        "coding_smoke_passed": coding_smoke_passed, "category_summary": category_summary,
        "screening_unprocessed_path": str(unprocessed_path), "categorized_candidate_csv": str(coding_csv),
        "categorized_candidate_workbook": str(workbook_path), "uncoded_survivor_path": str(uncoded_path),
        "screening_is_accuracy_validation": False, "coding_is_final_full_report_extraction": False,
    })
    atomic_json(output_dir / "run_report.json", snapshot)
    binding["status"] = snapshot["status"]
    atomic_json(output_dir / "terminal_status.json", {"status": snapshot["status"], "stopping_reason": snapshot["stopping_reason"], "completed_at_utc": utc_now()})
    fcntl.flock(lock_handle.fileno(), fcntl.LOCK_UN)
    lock_handle.close()
    return snapshot
