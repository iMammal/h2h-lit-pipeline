"""Staging-only high-precision re-screen of the frozen 890 INCLUDE candidates."""

from __future__ import annotations

import csv
import fcntl
import hashlib
import json
import os
import signal
import threading
import time
from collections import deque
from concurrent.futures import FIRST_COMPLETED, Future, ThreadPoolExecutor, wait
from pathlib import Path
from typing import Any, Iterable, Mapping

from h2h_lit.inference import InferenceProvider
from h2h_lit.openai_provider import OpenAIResponsesProvider
from h2h_lit.title_abstract_fast_track_batch import (
    CRITERION_KEYS,
    Pricing,
    RunSettings,
    _maximum_attempt_cost,
    _normalize_usage,
    _usage_cost,
    load_jsonl,
    sha256_file,
    validate_batch,
)
from h2h_lit.title_abstract_fast_track_overnight import (
    Budget,
    RuntimeWindow,
    _error_category,
    atomic_json,
    create_json,
    json_hash,
    stable_id,
    utc_now,
)
from h2h_lit.title_abstract_fast_track_repair import (
    EVIDENCED,
    NOT_EVIDENCED,
    _e7_yes_needs_consistency_warning,
    evidence_unit_response_schema,
    evidence_units_for,
    inference_input_for,
    validate_evidence_unit_payload,
)
from h2h_lit.title_abstract_screening import E6_STATUS

AMENDMENT_VERSION = "2.2.0"
PROMPT_VERSION = "2.2.0"
OUTPUT_SCHEMA_VERSION = "2.2.0"
RUNNER_VERSION = "1.0.0"
DOCUMENT_SCOPE_VALUES = {
    "QUALIFYING_RESEARCH_REPORT",
    "SURVEY_REVIEW_OVERVIEW",
    "UNCERTAIN",
}
WORKFLOW_INTEGRATION_VALUES = {
    "SUPPORTED",
    "CONTRADICTED",
    "NOT_EVIDENCED_IN_SUPPLIED_METADATA",
}
_USAGE_KEYS = (
    "input_tokens",
    "cached_input_tokens",
    "cache_write_tokens",
    "output_tokens",
    "reasoning_tokens",
)


def _policy_item_schema(values: list[str]) -> dict[str, Any]:
    return {
        "type": "object",
        "additionalProperties": False,
        "properties": {
            "decision": {"type": "string", "enum": values},
            "evidence_status": {
                "type": "string",
                "enum": [EVIDENCED, NOT_EVIDENCED],
            },
            "evidence_ids": {
                "type": "array",
                "items": {"type": "string"},
            },
            "rationale": {"type": "string", "minLength": 1},
        },
        "required": ["decision", "evidence_status", "evidence_ids", "rationale"],
    }


def precision_response_schema() -> dict[str, Any]:
    """Strict Responses API schema for the prospective v2.2.0 pass."""

    base = evidence_unit_response_schema()
    return {
        "type": "object",
        "additionalProperties": False,
        "properties": {
            "criteria": base["properties"]["criteria"],
            "policy_assessment": {
                "type": "object",
                "additionalProperties": False,
                "properties": {
                    "document_scope": _policy_item_schema(sorted(DOCUMENT_SCOPE_VALUES)),
                    "integrated_workflow": _policy_item_schema(
                        sorted(WORKFLOW_INTEGRATION_VALUES)
                    ),
                },
                "required": ["document_scope", "integrated_workflow"],
            },
            "overall_rationale": {"type": "string", "minLength": 1},
        },
        "required": ["criteria", "policy_assessment", "overall_rationale"],
    }


def _validate_policy_item(
    item: Any,
    *,
    field: str,
    allowed: set[str],
    evidence_units: Mapping[str, Mapping[str, Any]],
) -> dict[str, Any]:
    expected = {"decision", "evidence_status", "evidence_ids", "rationale"}
    if not isinstance(item, dict) or set(item) != expected:
        raise ValueError(f"policy_assessment.{field} has invalid fields")
    decision = item["decision"]
    if decision not in allowed:
        raise ValueError(f"policy_assessment.{field}.decision is invalid")
    status = item["evidence_status"]
    if status not in {EVIDENCED, NOT_EVIDENCED}:
        raise ValueError(f"policy_assessment.{field}.evidence_status is invalid")
    evidence_ids = item["evidence_ids"]
    if (
        not isinstance(evidence_ids, list)
        or any(not isinstance(value, str) for value in evidence_ids)
        or len(set(evidence_ids)) != len(evidence_ids)
    ):
        raise ValueError(f"policy_assessment.{field}.evidence_ids must be unique strings")
    unknown = sorted(set(evidence_ids) - set(evidence_units))
    if unknown:
        raise ValueError(
            f"policy_assessment.{field}.evidence_ids contain nonexistent IDs: {unknown}"
        )
    if status == EVIDENCED and not evidence_ids:
        raise ValueError(f"policy_assessment.{field}.EVIDENCED requires evidence IDs")
    if status == NOT_EVIDENCED and evidence_ids:
        raise ValueError(
            f"policy_assessment.{field}.NOT_EVIDENCED requires an empty evidence ID list"
        )
    evidenced_decisions = {
        "QUALIFYING_RESEARCH_REPORT",
        "SURVEY_REVIEW_OVERVIEW",
        "SUPPORTED",
        "CONTRADICTED",
    }
    if decision in evidenced_decisions and status != EVIDENCED:
        raise ValueError(f"policy_assessment.{field}.{decision} requires evidence")
    if decision == "NOT_EVIDENCED_IN_SUPPLIED_METADATA" and status != NOT_EVIDENCED:
        raise ValueError(
            "integrated_workflow.NOT_EVIDENCED_IN_SUPPLIED_METADATA requires absent-evidence status"
        )
    rationale = item["rationale"]
    if not isinstance(rationale, str) or not rationale.strip():
        raise ValueError(f"policy_assessment.{field}.rationale must be non-empty")
    return {
        **item,
        "rationale": rationale.strip(),
        "evidence": [dict(evidence_units[value]) for value in evidence_ids],
    }


def validate_precision_payload(payload: Any, record: Mapping[str, Any]) -> dict[str, Any]:
    """Validate evidence IDs and independently aggregate the v2.2.0 selection policy."""

    if not isinstance(payload, dict) or set(payload) != {
        "criteria",
        "policy_assessment",
        "overall_rationale",
    }:
        raise ValueError(
            "response must contain exactly criteria, policy_assessment, and overall_rationale"
        )
    base = validate_evidence_unit_payload(
        {
            "criteria": payload["criteria"],
            "overall_rationale": payload["overall_rationale"],
        },
        record,
    )
    policy = payload["policy_assessment"]
    if not isinstance(policy, dict) or set(policy) != {
        "document_scope",
        "integrated_workflow",
    }:
        raise ValueError("policy_assessment has invalid fields")
    units = {unit["id"]: unit for unit in evidence_units_for(record)}
    document_scope = _validate_policy_item(
        policy["document_scope"],
        field="document_scope",
        allowed=DOCUMENT_SCOPE_VALUES,
        evidence_units=units,
    )
    integrated_workflow = _validate_policy_item(
        policy["integrated_workflow"],
        field="integrated_workflow",
        allowed=WORKFLOW_INTEGRATION_VALUES,
        evidence_units=units,
    )

    responses = base["responses"]
    scientific_no = [
        key for key in ("E1", "E2", "E3", "E4", "E5") if responses[key] == "NO"
    ]
    scientific_uncertain = any(
        responses[key] == "UNCERTAIN" for key in ("E1", "E2", "E3", "E4", "E5")
    )
    abstract_missing = not str(record.get("abstract", "")).strip()
    scope = document_scope["decision"]
    integration = integrated_workflow["decision"]
    exclusion_reasons = list(base["exclusion_reasons"])

    if scope == "SURVEY_REVIEW_OVERVIEW":
        outcome = "EXCLUDED"
        disposition = "EXCLUDE"
        route = "H2H3_BACKGROUND"
        exclusion_reasons.append("SURVEY_REVIEW_OR_BROAD_OVERVIEW")
    elif scientific_no:
        outcome = "EXCLUDED"
        disposition = "EXCLUDE"
        route = "NONE"
    elif integration == "CONTRADICTED":
        outcome = "EXCLUDED"
        disposition = "EXCLUDE"
        route = "NONE"
        exclusion_reasons.append("INTEGRATED_ANALYTICAL_WORKFLOW_CONTRADICTED")
    elif (
        scope == "UNCERTAIN"
        or integration == "NOT_EVIDENCED_IN_SUPPLIED_METADATA"
        or scientific_uncertain
        or responses["E7"] != "YES"
        or abstract_missing
    ):
        outcome = "UNCERTAIN"
        disposition = "DEFER"
        route = "NONE"
    else:
        outcome = "INCLUDE"
        disposition = "ADVANCE_TO_FULL_REPORT_ASSESSMENT"
        route = "FULL_REPORT_ASSESSMENT"

    return {
        "criteria": base["criteria"],
        "policy_assessment": {
            "document_scope": document_scope,
            "integrated_workflow": integrated_workflow,
        },
        "overall_rationale": str(payload["overall_rationale"]).strip(),
        "responses": responses,
        "e6_status": E6_STATUS,
        "computed_outcome": outcome,
        "operational_disposition": disposition,
        "route": route,
        "exclusion_reasons": sorted(set(exclusion_reasons)),
    }


def result_group(result: Mapping[str, Any]) -> str:
    if result.get("status") == "AMBIGUOUS":
        return "ambiguous"
    if result.get("status") != "VALIDATED":
        return "unprocessed" if str(result.get("status", "")).startswith("UNPROCESSED") else "failed"
    return {
        "ADVANCE_TO_FULL_REPORT_ASSESSMENT": "advance",
        "DEFER": "deferred",
        "EXCLUDE": "excluded",
    }[result["judgment"]["operational_disposition"]]


def _evidence_summary(judgment: Mapping[str, Any]) -> str:
    rows = []
    for key in CRITERION_KEYS:
        item = judgment["criteria"][key]
        evidence = "; ".join(
            f"{unit['id']}: {unit['text']}" for unit in item.get("evidence", [])
        )
        rows.append(f"{key}={item['decision']} | {item['rationale']} | {evidence}")
    return "\n".join(rows)


def _result_csv_row(result: Mapping[str, Any], source: Mapping[str, Any]) -> dict[str, Any]:
    judgment = result.get("judgment", {})
    policy = judgment.get("policy_assessment", {})
    return {
        "batch_order": result.get("batch_order"),
        "prior_candidate_number": source.get("prior_candidate_number"),
        "canonical_id": result.get("canonical_id"),
        "title": source.get("title", ""),
        "abstract": source.get("abstract", ""),
        "abstract_status": "MISSING" if not str(source.get("abstract", "")).strip() else "PRESENT",
        "screening_status": result.get("status", "UNPROCESSED"),
        "doi": source.get("doi") or "",
        "source_database": source.get("source_database") or "",
        "source_identifier": source.get("source_identifier") or "",
        "source_url": source.get("source_url") or "",
        "outcome": judgment.get("computed_outcome", ""),
        "disposition": judgment.get("operational_disposition", ""),
        "route": judgment.get("route", ""),
        "document_scope": policy.get("document_scope", {}).get("decision", ""),
        "integrated_workflow": policy.get("integrated_workflow", {}).get("decision", ""),
        "exclusion_reasons": "; ".join(judgment.get("exclusion_reasons", [])),
        "criterion_evidence": _evidence_summary(judgment) if judgment else "",
        "overall_rationale": judgment.get("overall_rationale", ""),
        "e6_status": judgment.get("e6_status", E6_STATUS),
        "human_full_report_availability": "",
        "human_scientific_eligibility": "",
        "human_e6_verification": "",
        "final_paper_eligibility": "",
        "synthesis_code": "",
        "human_notes": "",
    }


CSV_FIELDS = list(
    _result_csv_row(
        {},
        {},
    )
)


class PrecisionProgress:
    def __init__(
        self,
        output_dir: Path,
        records: list[dict[str, Any]],
        initial_cost: float,
        initial_usage: Mapping[str, int],
    ):
        self.output_dir = output_dir
        self.records = records
        self.by_id = {row["canonical_id"]: row for row in records}
        self.initial_cost = initial_cost
        self.initial_usage = {key: int(initial_usage.get(key, 0)) for key in _USAGE_KEYS}
        self.counts = {
            name: 0
            for name in ("advance", "deferred", "excluded", "background", "failed", "ambiguous")
        }
        self.usage = {key: 0 for key in _USAGE_KEYS}
        self.cost = 0.0
        self.latency = 0.0
        self.requests = 0
        self.retries = 0
        self.warnings = 0
        self.completed_ids: set[str] = set()
        self.lock = threading.Lock()
        self.events_path = output_dir / "completion_events.jsonl"
        self.csv_paths = {
            "advance": output_dir / "advance_candidates.csv",
            "deferred": output_dir / "deferred.csv",
            "excluded": output_dir / "excluded.csv",
            "background": output_dir / "h2h3_background.csv",
        }
        for target in self.csv_paths.values():
            if not target.exists():
                with target.open("x", encoding="utf-8", newline="") as handle:
                    csv.DictWriter(handle, fieldnames=CSV_FIELDS).writeheader()

    def add(self, result: Mapping[str, Any], *, persist_event: bool = True) -> None:
        record_id = str(result["canonical_id"])
        with self.lock:
            if record_id in self.completed_ids:
                return
            self.completed_ids.add(record_id)
            group = result_group(result)
            if group != "unprocessed":
                self.counts[group] += 1
            is_background = (
                result.get("status") == "VALIDATED"
                and result.get("judgment", {}).get("route") == "H2H3_BACKGROUND"
            )
            if is_background:
                self.counts["background"] += 1
            for attempt in result.get("attempts", []):
                self.requests += 1
                self.retries += int(int(attempt.get("attempt_number", 1)) > 1)
                self.cost += float(attempt.get("estimated_billed_cost_usd", 0.0))
                self.latency += float(attempt.get("latency_seconds", 0.0))
                for key in _USAGE_KEYS:
                    self.usage[key] += int(attempt.get("usage", {}).get(key, 0))
            if result.get("e7_consistency_warning"):
                self.warnings += 1
            if not persist_event:
                return
            with self.events_path.open("a", encoding="utf-8", newline="\n") as handle:
                handle.write(json.dumps({
                    "completed_at_utc": utc_now(),
                    "batch_order": result.get("batch_order"),
                    "canonical_id": record_id,
                    "status": result.get("status"),
                    "group": group,
                    "background": is_background,
                }, sort_keys=True) + "\n")
            if group in self.csv_paths:
                with self.csv_paths[group].open("a", encoding="utf-8", newline="") as handle:
                    csv.DictWriter(handle, fieldnames=CSV_FIELDS).writerow(
                        _result_csv_row(result, self.by_id[record_id])
                    )
            if is_background:
                with self.csv_paths["background"].open("a", encoding="utf-8", newline="") as handle:
                    csv.DictWriter(handle, fieldnames=CSV_FIELDS).writerow(
                        _result_csv_row(result, self.by_id[record_id])
                    )

    def snapshot(
        self,
        budget: Budget,
        runtime: RuntimeWindow,
        stop_reason: str | None,
    ) -> dict[str, Any]:
        with self.lock:
            budget_state = budget.snapshot()
            processed = len(self.completed_ids)
            return {
                "updated_at_utc": utc_now(),
                "status": "RUNNING" if stop_reason is None else "STOPPING",
                "stop_reason": stop_reason,
                "batch_records": len(self.records),
                "processed_records": processed,
                "valid_screenings": (
                    self.counts["advance"] + self.counts["deferred"] + self.counts["excluded"]
                ),
                "unprocessed_records": len(self.records) - processed,
                "counts": dict(self.counts),
                "new_run_usage": dict(self.usage),
                "cumulative_usage": {
                    key: self.initial_usage[key] + self.usage[key] for key in _USAGE_KEYS
                },
                "new_run_estimated_cost_usd": self.cost,
                "cumulative_conservative_cost_usd": budget_state["actual"],
                "current_inflight_reserved_cost_usd": budget_state["reserved"],
                "hard_cumulative_spending_cap_usd": budget_state["cap"],
                "requests": self.requests,
                "retries": self.retries,
                "aggregate_latency_seconds": self.latency,
                "e7_consistency_warnings": self.warnings,
                "elapsed_seconds_from_first_new_request": runtime.elapsed(),
                "throughput_records_per_hour": (
                    processed / runtime.elapsed() * 3600.0 if runtime.elapsed() > 0 else 0.0
                ),
            }


def _screen_record(
    record: dict[str, Any],
    records_dir: Path,
    prompt: str,
    prompt_hash: str,
    settings: RunSettings,
    provider: InferenceProvider,
    budget: Budget,
    runtime: RuntimeWindow,
) -> dict[str, Any]:
    record_dir = records_dir / record["canonical_id"].replace(":", "_")
    record_dir.mkdir(exist_ok=True)
    final_path = record_dir / "result.json"
    if final_path.exists():
        return json.loads(final_path.read_text(encoding="utf-8"))
    input_snapshot = inference_input_for(record)
    input_hash = json_hash(input_snapshot)
    parameters = {
        "max_output_tokens": settings.max_output_tokens,
        "reasoning_effort": settings.reasoning_effort,
        "response_schema_version": OUTPUT_SCHEMA_VERSION,
        "service_tier": settings.service_tier,
        "store": False,
        "structured_output": "title_abstract_precision_rescreen_v2_2_0",
        "verbosity": settings.verbosity,
    }
    attempts: list[dict[str, Any]] = []
    for attempt_number in range(1, settings.retry_limit + 2):
        reservation_path = record_dir / f"attempt-{attempt_number:03d}.reservation.json"
        attempt_path = record_dir / f"attempt-{attempt_number:03d}.json"
        if reservation_path.exists() and not attempt_path.exists():
            result = {
                "status": "AMBIGUOUS",
                "failure_reason": "OUTSTANDING_RESERVATION_NOT_RESUBMITTED",
                "canonical_id": record["canonical_id"],
                "batch_order": record["batch_order"],
                "title": record["title"],
                "attempts": attempts,
            }
            create_json(final_path, result)
            return result
        if not runtime.allow_new_request():
            status = "UNPROCESSED_RUNTIME_STOP" if not attempts else "FAILED"
            result = {
                "status": status,
                "failure_reason": "RUNTIME_DRAIN_STOP_BEFORE_REQUEST",
                "canonical_id": record["canonical_id"],
                "batch_order": record["batch_order"],
                "title": record["title"],
                "attempts": attempts,
            }
            create_json(final_path, result)
            return result
        maximum = _maximum_attempt_cost(prompt, input_snapshot, settings)
        if not budget.reserve(maximum):
            status = "UNPROCESSED_BUDGET_STOP" if not attempts else "FAILED"
            result = {
                "status": status,
                "failure_reason": "CUMULATIVE_BUDGET_STOP_BEFORE_REQUEST",
                "canonical_id": record["canonical_id"],
                "batch_order": record["batch_order"],
                "title": record["title"],
                "attempts": attempts,
            }
            create_json(final_path, result)
            return result
        runtime.mark_first_request()
        request_id = stable_id(
            "precision-rescreen-v2.2.0",
            record["canonical_id"],
            input_hash,
            prompt_hash,
            str(attempt_number),
        )
        create_json(reservation_path, {
            "request_id": request_id,
            "attempt_number": attempt_number,
            "maximum_reserved_cost_usd": maximum,
            "input_hash": input_hash,
        })
        started = time.monotonic()
        raw_response = ""
        metadata: dict[str, Any] = {}
        validated = None
        error = None
        try:
            raw_response = provider.generate(
                model=settings.model,
                prompt=prompt,
                input_snapshot=input_snapshot,
                parameters=parameters,
                request_id=request_id,
                attempt_number=attempt_number,
            )
            validated = validate_precision_payload(json.loads(raw_response), record)
        except Exception as exc:  # noqa: BLE001 - persisted as run evidence
            error = f"{type(exc).__name__}: {exc}"
        try:
            if isinstance(provider, OpenAIResponsesProvider) or hasattr(provider, "metadata_for"):
                metadata = dict(provider.metadata_for(request_id, attempt_number))
        except Exception as exc:  # noqa: BLE001
            error = error or f"metadata error: {type(exc).__name__}: {exc}"
            validated = None
        try:
            usage = _normalize_usage(metadata.get("provider_usage"))
        except (TypeError, ValueError) as exc:
            usage = {key: 0 for key in _USAGE_KEYS}
            error = error or f"invalid provider usage: {exc}"
            validated = None
        actual = _usage_cost(usage, settings.pricing) if any(usage.values()) else maximum
        budget.settle(maximum, actual)
        attempt = {
            "request_id": request_id,
            "attempt_number": attempt_number,
            "status": "VALID" if validated is not None else "INVALID",
            "latency_seconds": max(0.0, time.monotonic() - started),
            "usage": usage,
            "estimated_billed_cost_usd": actual,
            "raw_response": raw_response,
            "validation_error": error,
            "provider_metadata": metadata,
        }
        create_json(attempt_path, attempt)
        attempts.append(attempt)
        if validated is not None:
            warning = _e7_yes_needs_consistency_warning(validated["responses"])
            result = {
                "status": "VALIDATED",
                "canonical_id": record["canonical_id"],
                "batch_order": record["batch_order"],
                "prior_candidate_number": record["prior_candidate_number"],
                "title": record["title"],
                "abstract_missing": not bool(str(record["abstract"]).strip()),
                "doi": record.get("doi"),
                "source_url": record.get("source_url"),
                "source_database": record.get("source_database"),
                "source_identifier": record.get("source_identifier"),
                "input_hash": input_hash,
                "prompt_hash": prompt_hash,
                "output_schema_version": OUTPUT_SCHEMA_VERSION,
                "judgment": validated,
                "e7_consistency_warning": warning,
                "attempts": attempts,
            }
            create_json(final_path, result)
            return result
        category = _error_category({"attempts": attempts})
        if category in {"AUTHENTICATION", "BILLING"}:
            break
    result = {
        "status": "FAILED",
        "failure_reason": "RETRIES_EXHAUSTED",
        "canonical_id": record["canonical_id"],
        "batch_order": record["batch_order"],
        "prior_candidate_number": record["prior_candidate_number"],
        "title": record["title"],
        "input_hash": input_hash,
        "prompt_hash": prompt_hash,
        "attempts": attempts,
    }
    create_json(final_path, result)
    return result


def _load_existing_results(records_dir: Path) -> dict[str, dict[str, Any]]:
    results = {}
    for path in records_dir.glob("*/result.json"):
        value = json.loads(path.read_text(encoding="utf-8"))
        record_id = value.get("canonical_id")
        if not isinstance(record_id, str) or record_id in results:
            raise ValueError("invalid or duplicate persisted precision result")
        results[record_id] = value
    return results


def _write_jsonl(path: Path, rows: Iterable[Mapping[str, Any]]) -> None:
    with path.open("w", encoding="utf-8", newline="\n") as handle:
        for row in rows:
            handle.write(json.dumps(row, sort_keys=True, ensure_ascii=True) + "\n")


def _write_simple_csv(
    path: Path,
    rows: Iterable[Mapping[str, Any]],
    records_by_id: Mapping[str, Mapping[str, Any]],
) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=CSV_FIELDS)
        writer.writeheader()
        for row in rows:
            source = records_by_id[str(row["canonical_id"])]
            writer.writerow(_result_csv_row(row, source))


def _write_transition_table(
    path: Path,
    records: list[dict[str, Any]],
    results: Mapping[str, Mapping[str, Any]],
) -> None:
    fields = [
        "batch_order",
        "prior_candidate_number",
        "canonical_id",
        "title",
        "prior_outcome",
        "new_status",
        "new_outcome",
        "new_disposition",
        "route",
        "reason",
    ]
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for record in records:
            result = results.get(record["canonical_id"])
            judgment = result.get("judgment", {}) if result else {}
            writer.writerow({
                "batch_order": record["batch_order"],
                "prior_candidate_number": record["prior_candidate_number"],
                "canonical_id": record["canonical_id"],
                "title": record["title"],
                "prior_outcome": "INCLUDE",
                "new_status": result.get("status", "UNPROCESSED") if result else "UNPROCESSED",
                "new_outcome": judgment.get("computed_outcome", ""),
                "new_disposition": judgment.get("operational_disposition", ""),
                "route": judgment.get("route", ""),
                "reason": "; ".join(judgment.get("exclusion_reasons", []))
                or judgment.get("overall_rationale", ""),
            })


def _write_advance_workbook(
    path: Path,
    advance_rows: list[Mapping[str, Any]],
    records_by_id: Mapping[str, Mapping[str, Any]],
    transition_path: Path,
) -> None:
    """Write the final editable handoff after screening completes or stops."""

    from openpyxl import Workbook  # deployment-pinned staging dependency
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    workbook = Workbook()
    instructions = workbook.active
    instructions.title = "Instructions"
    instructions.append(["High-precision candidate handoff"])
    instructions.append([
        "Model-nominated candidates requiring human full-report assessment. "
        "These are not final eligible papers or an independent accuracy benchmark."
    ])
    instructions.append(["E6 remains NOT_ASSESSED_AT_THIS_STAGE."])
    instructions.append(["Historical INCLUDE outcomes remain preserved in the transition table."])
    queue = workbook.create_sheet("ADVANCE candidates")
    queue_fields = CSV_FIELDS
    queue.append(queue_fields)
    for result in advance_rows:
        values = _result_csv_row(result, records_by_id[str(result["canonical_id"])])
        queue.append([values[field] for field in queue_fields])
    transitions = workbook.create_sheet("Transitions")
    with transition_path.open(encoding="utf-8", newline="") as handle:
        for row in csv.reader(handle):
            transitions.append(row)

    dark = PatternFill("solid", fgColor="263B50")
    input_fill = PatternFill("solid", fgColor="FFF4CC")
    for sheet in (queue, transitions):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = dark
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    human_columns = [
        queue_fields.index("human_full_report_availability") + 1,
        queue_fields.index("human_scientific_eligibility") + 1,
        queue_fields.index("human_e6_verification") + 1,
        queue_fields.index("final_paper_eligibility") + 1,
        queue_fields.index("synthesis_code") + 1,
        queue_fields.index("human_notes") + 1,
    ]
    for column in human_columns:
        for row in range(2, queue.max_row + 1):
            queue.cell(row=row, column=column).fill = input_fill
    validations = [
        ("human_full_report_availability", '"AVAILABLE,UNAVAILABLE,NOT_CHECKED"'),
        ("human_scientific_eligibility", '"ELIGIBLE,NOT_ELIGIBLE,UNCERTAIN"'),
        ("human_e6_verification", '"VERIFIED,FAILED,NOT_ASSESSED"'),
        ("final_paper_eligibility", '"INCLUDE,EXCLUDE,DEFER,NOT_ASSESSED"'),
    ]
    for field, formula in validations:
        column = queue_fields.index(field) + 1
        validation = DataValidation(type="list", formula1=formula, allow_blank=True)
        queue.add_data_validation(validation)
        validation.add(f"{queue.cell(2, column).coordinate}:{queue.cell(max(2, queue.max_row), column).coordinate}")
    widths = {
        "A": 12, "B": 15, "C": 30, "D": 55, "E": 80, "F": 14,
        "G": 18, "H": 20, "I": 20, "J": 24, "K": 42, "L": 14,
        "M": 20, "N": 16, "O": 24, "P": 24, "Q": 24, "R": 90,
        "S": 55, "T": 30, "U": 25, "V": 24, "W": 24, "X": 22,
        "Y": 20, "Z": 45,
    }
    for column, width in widths.items():
        queue.column_dimensions[column].width = width
    queue.row_dimensions[1].height = 36
    instructions.column_dimensions["A"].width = 110
    for row in range(1, 5):
        instructions.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="top")
    workbook.save(path)


def run_precision_rescreen(
    *,
    batch_path: Path,
    batch_manifest_path: Path,
    prompt_path: Path,
    protocol_path: Path,
    amendment_path: Path,
    campaign_ledger_path: Path,
    output_dir: Path,
    settings: RunSettings,
    provider: InferenceProvider,
    runtime_hours: float = 3.0,
    drain_seconds: float = 300.0,
    resume: bool = False,
) -> dict[str, Any]:
    settings.validate()
    if settings.smoke_count != 20 or settings.concurrency != 4 or settings.retry_limit > 1:
        raise ValueError("precision re-screen requires smoke=20, concurrency=4, retry_limit<=1")
    records = list(load_jsonl(batch_path))
    validate_batch(records, expected_count=890)
    manifest = json.loads(batch_manifest_path.read_text(encoding="utf-8"))
    if (
        manifest.get("batch", {}).get("sha256") != sha256_file(batch_path)
        or manifest.get("batch", {}).get("records") != 890
        or manifest.get("batch", {}).get("unique_ids") != 890
    ):
        raise ValueError("precision batch binding mismatch")
    if manifest.get("counts") != {
        "records": 890,
        "prior_250_effective_repair": 11,
        "overnight_25000": 879,
        "known_related_version_candidates": 5,
    }:
        raise ValueError("precision cohort binding drift")
    if manifest.get("campaign_ledger", {}).get("sha256") != sha256_file(campaign_ledger_path):
        raise ValueError("campaign spending ledger binding mismatch")
    protocol = json.loads(protocol_path.read_text(encoding="utf-8"))
    amendment = json.loads(amendment_path.read_text(encoding="utf-8"))
    if protocol.get("protocol_version") != "2.0.0":
        raise ValueError("unexpected base protocol version")
    if amendment.get("amendment_version") != AMENDMENT_VERSION:
        raise ValueError("unexpected precision amendment version")
    if manifest.get("protocol", {}).get("sha256") != sha256_file(protocol_path):
        raise ValueError("protocol manifest binding mismatch")
    if manifest.get("amendment", {}).get("sha256") != sha256_file(amendment_path):
        raise ValueError("amendment manifest binding mismatch")
    prompt = prompt_path.read_text(encoding="utf-8")
    if (
        f"Prompt-Version: {PROMPT_VERSION}" not in prompt
        or f"Output-Schema-Version: {OUTPUT_SCHEMA_VERSION}" not in prompt
    ):
        raise ValueError("unexpected precision prompt/schema version")
    prompt_hash = hashlib.sha256(prompt.encode("utf-8")).hexdigest()
    if manifest.get("prompt", {}).get("sha256") != prompt_hash:
        raise ValueError("prompt manifest binding mismatch")
    campaign_ledger = json.loads(campaign_ledger_path.read_text(encoding="utf-8"))
    initial_cost = float(campaign_ledger["cumulative_conservative_cost_usd"])
    unresolved_reservation = float(campaign_ledger.get("current_inflight_reserved_cost_usd", 0.0))
    initial_budget_cost = initial_cost + max(0.0, unresolved_reservation)
    initial_usage = campaign_ledger["cumulative_usage"]
    if initial_budget_cost > settings.hard_spending_cap_usd:
        raise ValueError("persisted campaign cost already exceeds the authorized cap")

    if output_dir.exists() and not resume:
        raise FileExistsError("precision output directory exists; explicit resume required")
    output_dir.mkdir(parents=True, exist_ok=True)
    records_dir = output_dir / "records"
    records_dir.mkdir(exist_ok=True)
    lock_handle = (output_dir / "budget_owner.lock").open("a+", encoding="utf-8")
    try:
        fcntl.flock(lock_handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
    except BlockingIOError as exc:
        raise RuntimeError("another worker owns the precision campaign ledger") from exc

    binding = {
        "artifact_class": "title_abstract_precision_rescreen_run",
        "runner_version": RUNNER_VERSION,
        "status": "RUNNING",
        "created_at_utc": utc_now(),
        "batch": {"path": str(batch_path), "sha256": sha256_file(batch_path), "records": 890},
        "batch_manifest": {"path": str(batch_manifest_path), "sha256": sha256_file(batch_manifest_path)},
        "source_bindings": manifest["source_bindings"],
        "protocol": {"path": str(protocol_path), "sha256": sha256_file(protocol_path), "version": "2.0.0"},
        "amendment": {"path": str(amendment_path), "sha256": sha256_file(amendment_path), "version": AMENDMENT_VERSION},
        "prompt": {"path": str(prompt_path), "sha256": prompt_hash, "version": PROMPT_VERSION},
        "schema": {"version": OUTPUT_SCHEMA_VERSION, "sha256": json_hash(precision_response_schema())},
        "implementation": {"path": str(Path(__file__)), "sha256": sha256_file(Path(__file__))},
        "campaign_ledger": {
            "path": str(campaign_ledger_path),
            "sha256": sha256_file(campaign_ledger_path),
            "starting_cumulative_conservative_cost_usd": initial_cost,
            "carried_unresolved_reservations_usd": unresolved_reservation,
        },
        "provider": settings.provider_name,
        "model": settings.model,
        "parameters": {
            "reasoning_effort": settings.reasoning_effort,
            "service_tier": settings.service_tier,
            "max_output_tokens": settings.max_output_tokens,
            "verbosity": settings.verbosity,
            "store": False,
            "tools": [],
        },
        "pricing_usd_per_million_tokens": {
            "input": settings.pricing.input_usd_per_million,
            "cached_input": settings.pricing.cached_input_usd_per_million,
            "cache_write": settings.pricing.cache_write_usd_per_million,
            "output": settings.pricing.output_usd_per_million,
        },
        "controls": {
            "hard_cumulative_spending_cap_usd": settings.hard_spending_cap_usd,
            "runtime_hours": runtime_hours,
            "drain_seconds": drain_seconds,
            "record_limit": 890,
            "smoke_count": 20,
            "concurrency": 4,
            "retry_limit": settings.retry_limit,
            "timeout_seconds": settings.timeout_seconds,
            "rolling_failure_gate": "3 retry-exhausted records in 20 completions",
        },
        "model_input_policy": "canonical_id and deterministic original title/abstract evidence units only",
        "prior_outcomes_exposed_to_model": False,
        "staging_only": True,
    }
    binding_path = output_dir / "run_binding.json"
    if binding_path.exists():
        existing_binding = json.loads(binding_path.read_text(encoding="utf-8"))
        for key in (
            "batch", "prompt", "schema", "model", "parameters",
            "pricing_usd_per_million_tokens", "controls", "campaign_ledger",
        ):
            if existing_binding.get(key) != binding.get(key):
                raise ValueError(f"resume binding mismatch: {key}")
        binding = existing_binding
    else:
        create_json(binding_path, binding)

    existing = _load_existing_results(records_dir)
    persisted_cost = sum(
        float(attempt.get("estimated_billed_cost_usd", 0.0))
        for result in existing.values()
        for attempt in result.get("attempts", [])
    )
    budget = Budget(settings.hard_spending_cap_usd, initial_budget_cost + persisted_cost)
    runtime = RuntimeWindow(output_dir, runtime_hours, drain_seconds)
    progress = PrecisionProgress(output_dir, records, initial_budget_cost, initial_usage)
    for result in existing.values():
        progress.add(result, persist_event=False)

    stop_event = threading.Event()
    stop_reason: str | None = None
    rolling = deque(maxlen=20)
    persistent_categories = deque(maxlen=3)

    def request_stop(reason: str) -> None:
        nonlocal stop_reason
        if stop_reason is None:
            stop_reason = reason
            stop_event.set()
            print(json.dumps({"event": "stop_requested", "reason": reason, "at": utc_now()}), flush=True)

    def signal_stop(signum: int, _frame: Any) -> None:
        request_stop(f"SIGNAL_{signum}_GRACEFUL_DRAIN")

    previous_handlers = {
        signum: signal.signal(signum, signal_stop) for signum in (signal.SIGTERM, signal.SIGINT)
    }

    def register(result: dict[str, Any]) -> None:
        progress.add(result)
        rolling.append(
            result.get("status") == "FAILED"
            and result.get("failure_reason") == "RETRIES_EXHAUSTED"
        )
        category = _error_category(result)
        if category in {"AUTHENTICATION", "BILLING"}:
            request_stop(f"{category}_FAILURE")
        if category in {"RATE_LIMIT", "STORAGE"}:
            persistent_categories.append(category)
            if len(persistent_categories) == 3 and len(set(persistent_categories)) == 1:
                request_stop(f"PERSISTENT_{category}_FAILURE")
        else:
            persistent_categories.clear()
        if len(rolling) == 20 and sum(rolling) >= 3:
            request_stop("ROLLING_VALIDATION_FAILURE_GATE")
        atomic_json(output_dir / "progress.json", progress.snapshot(budget, runtime, stop_reason))

    pending = [record for record in records if record["canonical_id"] not in existing]

    def run_phase(phase_records: list[dict[str, Any]], *, smoke: bool) -> list[dict[str, Any]]:
        iterator = iter(phase_records)
        futures: dict[Future[dict[str, Any]], dict[str, Any]] = {}
        completed: list[dict[str, Any]] = []
        with ThreadPoolExecutor(max_workers=min(4, max(1, len(phase_records)))) as pool:
            while True:
                while not stop_event.is_set() and len(futures) < 4:
                    if not runtime.allow_new_request():
                        request_stop("RUNTIME_LIMIT_DRAIN")
                        break
                    try:
                        record = next(iterator)
                    except StopIteration:
                        break
                    future = pool.submit(
                        _screen_record,
                        record,
                        records_dir,
                        prompt,
                        prompt_hash,
                        settings,
                        provider,
                        budget,
                        runtime,
                    )
                    futures[future] = record
                if not futures:
                    break
                done, _ = wait(futures, return_when=FIRST_COMPLETED)
                for future in done:
                    futures.pop(future)
                    result = future.result()
                    completed.append(result)
                    register(result)
                    print(json.dumps({
                        "event": "record_completed",
                        "batch_order": result.get("batch_order"),
                        "canonical_id": result.get("canonical_id"),
                        "status": result.get("status"),
                        "group": result_group(result),
                    }), flush=True)
                    if smoke and result.get("status") != "VALIDATED":
                        request_stop("SMOKE_TEST_VALIDATION_FAILURE")
        return completed

    smoke_passed = False
    try:
        smoke_records = pending[: settings.smoke_count]
        smoke_results = run_phase(smoke_records, smoke=True)
        smoke_passed = (
            len(smoke_results) == len(smoke_records) == settings.smoke_count
            and all(
                result.get("status") == "VALIDATED"
                and result.get("judgment", {}).get("computed_outcome")
                in {"INCLUDE", "UNCERTAIN", "EXCLUDED"}
                and result.get("judgment", {}).get("operational_disposition")
                in {"ADVANCE_TO_FULL_REPORT_ASSESSMENT", "DEFER", "EXCLUDE"}
                for result in smoke_results
            )
            and stop_reason is None
        )
        if smoke_passed:
            print(json.dumps({"event": "smoke_passed", "records": 20}), flush=True)
            run_phase(pending[settings.smoke_count:], smoke=False)
        elif stop_reason is None:
            request_stop("SMOKE_TEST_INCOMPLETE")
    finally:
        for signum, handler in previous_handlers.items():
            signal.signal(signum, handler)

    final_results = _load_existing_results(records_dir)
    groups = {name: [] for name in ("advance", "deferred", "excluded", "failed", "ambiguous", "unprocessed")}
    background = []
    for record in records:
        result = final_results.get(record["canonical_id"])
        if result is None:
            groups["unprocessed"].append(record)
            continue
        group = result_group(result)
        groups[group].append(result)
        if result.get("judgment", {}).get("route") == "H2H3_BACKGROUND":
            background.append(result)
    for name, rows in groups.items():
        _write_jsonl(output_dir / f"{name}.jsonl", rows)
    _write_jsonl(output_dir / "h2h3_background.jsonl", background)
    records_by_id = {record["canonical_id"]: record for record in records}
    for name in ("advance", "deferred", "excluded"):
        _write_simple_csv(output_dir / f"{name}.csv", groups[name], records_by_id)
    _write_simple_csv(output_dir / "h2h3_background.csv", background, records_by_id)
    for name in ("failed", "ambiguous"):
        _write_simple_csv(output_dir / f"{name}.csv", groups[name], records_by_id)
    _write_simple_csv(output_dir / "unprocessed.csv", groups["unprocessed"], records_by_id)
    transition_path = output_dir / "prior_include_to_precision_outcome.csv"
    _write_transition_table(transition_path, records, final_results)
    workbook_path = output_dir / "advance_candidates.xlsx"
    _write_advance_workbook(workbook_path, groups["advance"], records_by_id, transition_path)

    snapshot = progress.snapshot(budget, runtime, stop_reason)
    snapshot.update({
        "status": (
            "COMPLETE"
            if not groups["unprocessed"] and not groups["failed"] and not groups["ambiguous"]
            else "STOPPED"
        ),
        "stopping_reason": stop_reason or (
            "BATCH_COMPLETE"
            if not groups["unprocessed"] and not groups["failed"] and not groups["ambiguous"]
            else "BATCH_COMPLETE_WITH_FAILURES"
        ),
        "smoke_test_passed": smoke_passed,
        "counts": {
            "advance": len(groups["advance"]),
            "deferred": len(groups["deferred"]),
            "excluded": len(groups["excluded"]),
            "background": len(background),
            "failed": len(groups["failed"]),
            "ambiguous": len(groups["ambiguous"]),
            "unprocessed": len(groups["unprocessed"]),
        },
        "historical_prior_outcomes_preserved": True,
        "model_nominations_are_final_inclusions": False,
        "accuracy_assessed": False,
        "advance_candidate_csv": str(output_dir / "advance_candidates.csv"),
        "advance_candidate_workbook": str(workbook_path),
    })
    atomic_json(output_dir / "run_report.json", snapshot)
    (output_dir / "run_report.md").write_text(
        "# High-precision candidate re-screen\n\n"
        f"Status: `{snapshot['status']}`\n\n"
        f"- Stop reason: `{snapshot['stopping_reason']}`.\n"
        f"- ADVANCE: {len(groups['advance'])}; DEFER: {len(groups['deferred'])}; "
        f"EXCLUDED: {len(groups['excluded'])}; H2H3/background: {len(background)}.\n"
        f"- Failed: {len(groups['failed'])}; ambiguous: {len(groups['ambiguous'])}; "
        f"unprocessed: {len(groups['unprocessed'])}.\n"
        f"- Cumulative conservative cost: USD {snapshot['cumulative_conservative_cost_usd']:.6f} "
        f"of USD {snapshot['hard_cumulative_spending_cap_usd']:.2f}.\n"
        f"- Candidate CSV: `{output_dir / 'advance_candidates.csv'}`.\n"
        f"- Candidate workbook: `{workbook_path}`.\n\n"
        "Results are prospective model nominations requiring human full-report assessment.\n",
        encoding="utf-8",
        newline="\n",
    )
    artifacts = {}
    for artifact in sorted(output_dir.iterdir()):
        if artifact.is_file() and artifact.name not in {"package_manifest.json", "budget_owner.lock"}:
            artifacts[artifact.name] = {
                "sha256": sha256_file(artifact),
                "bytes": artifact.stat().st_size,
            }
    atomic_json(output_dir / "package_manifest.json", {
        "artifact_class": "title_abstract_precision_rescreen_staging_package",
        "runner_version": RUNNER_VERSION,
        "status": snapshot["status"],
        "counts": snapshot["counts"],
        "artifacts": artifacts,
        "staging_only": True,
    })
    fcntl.flock(lock_handle.fileno(), fcntl.LOCK_UN)
    lock_handle.close()
    return snapshot
