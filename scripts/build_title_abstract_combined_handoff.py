"""Build the terminal combined handoff for the remaining-corpus campaign."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
from pathlib import Path
from typing import Any, Iterable


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_jsonl(path: Path) -> Iterable[dict[str, Any]]:
    with path.open(encoding="utf-8") as handle:
        for line in handle:
            if line.strip():
                yield json.loads(line)


def create_json(path: Path, value: Any) -> None:
    payload = (json.dumps(value, indent=2, sort_keys=True) + "\n").encode("utf-8")
    descriptor = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
    with os.fdopen(descriptor, "wb") as handle:
        handle.write(payload)


def combine_csv(paths: list[Path], target: Path) -> tuple[int, set[str]]:
    seen: set[str] = set()
    count = 0
    fields: list[str] | None = None
    with target.open("x", encoding="utf-8", newline="") as output:
        writer = None
        for path in paths:
            with path.open(encoding="utf-8", newline="") as source:
                reader = csv.DictReader(source)
                if fields is None:
                    fields = list(reader.fieldnames or [])
                    writer = csv.DictWriter(output, fieldnames=fields)
                    writer.writeheader()
                elif list(reader.fieldnames or []) != fields:
                    raise ValueError(f"CSV schema mismatch: {path}")
                assert writer is not None
                for row in reader:
                    record_id = row.get("canonical_id", "")
                    if not record_id:
                        raise ValueError(f"missing canonical ID: {path}")
                    if record_id in seen:
                        raise ValueError(f"duplicate canonical ID across combined CSV: {record_id}")
                    seen.add(record_id)
                    writer.writerow(row)
                    count += 1
    return count, seen


def combine_categorized(
    *,
    parent_csv: Path,
    continuation_csv: Path,
    initial_ids: set[str],
    target: Path,
) -> tuple[int, set[str], list[str]]:
    seen: set[str] = set()
    fields: list[str] | None = None
    count = 0
    with target.open("x", encoding="utf-8", newline="") as output:
        writer = None
        for stage, path in (("PARENT_RUN", parent_csv), ("CONTINUATION", continuation_csv)):
            with path.open(encoding="utf-8", newline="") as source:
                reader = csv.DictReader(source)
                source_fields = list(reader.fieldnames or [])
                if fields is None:
                    fields = ["candidate_origin", "coding_stage_run", *source_fields]
                    writer = csv.DictWriter(output, fieldnames=fields)
                    writer.writeheader()
                elif source_fields != fields[2:]:
                    raise ValueError(f"categorized CSV schema mismatch: {path}")
                assert writer is not None
                for row in reader:
                    record_id = row.get("canonical_id", "")
                    if not record_id or record_id in seen:
                        raise ValueError(f"missing or duplicate categorized ID: {record_id}")
                    seen.add(record_id)
                    writer.writerow({
                        "candidate_origin": (
                            "INITIAL_PRECISION_ADVANCE_646"
                            if record_id in initial_ids
                            else "REMAINING_QUEUE_ADVANCE"
                        ),
                        "coding_stage_run": stage,
                        **row,
                    })
                    count += 1
    return count, seen, fields or []


def e7_warning_summary(result_dirs: list[Path]) -> dict[str, Any]:
    summary = {
        name: {"valid_screenings": 0, "e7_warnings": 0}
        for name in ("ADVANCE", "DEFER", "EXCLUDED")
    }
    seen: set[str] = set()
    for directory in result_dirs:
        for result_path in directory.glob("*/result.json"):
            result = json.loads(result_path.read_text(encoding="utf-8"))
            record_id = result.get("canonical_id")
            if not isinstance(record_id, str) or record_id in seen:
                raise ValueError(f"duplicate or malformed screening result: {result_path}")
            seen.add(record_id)
            if result.get("status") != "VALIDATED":
                continue
            disposition = result.get("judgment", {}).get("operational_disposition")
            group = {
                "ADVANCE_TO_FULL_REPORT_ASSESSMENT": "ADVANCE",
                "DEFER": "DEFER",
                "EXCLUDE": "EXCLUDED",
            }.get(disposition)
            if group is None:
                raise ValueError(f"invalid deterministic disposition: {record_id}")
            summary[group]["valid_screenings"] += 1
            summary[group]["e7_warnings"] += int(bool(result.get("e7_consistency_warning")))
    outcome_rows = tuple(summary.values())
    summary["total_valid_screenings"] = sum(row["valid_screenings"] for row in outcome_rows)
    summary["total_e7_warnings"] = sum(row["e7_warnings"] for row in outcome_rows)
    summary["advance_warning_check"] = (
        f"{summary['ADVANCE']['e7_warnings']} of "
        f"{summary['ADVANCE']['valid_screenings']} ADVANCE records carry a nonfatal E7 warning"
    )
    return summary


def sum_category_summaries(left: dict[str, Any], right: dict[str, Any]) -> dict[str, Any]:
    labels = {}
    for key in set(left.get("labels", {})) | set(right.get("labels", {})):
        labels[key] = int(left.get("labels", {}).get(key, 0)) + int(right.get("labels", {}).get(key, 0))
    workflow = {}
    for key in ("SUPPORTED", "UNSUPPORTED", "UNKNOWN"):
        workflow[key] = int(left.get("workflow_support_review", {}).get(key, 0)) + int(right.get("workflow_support_review", {}).get(key, 0))
    return {
        "unique_coded_papers": int(left.get("unique_coded_papers", 0)) + int(right.get("unique_coded_papers", 0)),
        "multilabel_present_assignments": int(left.get("multilabel_present_assignments", 0)) + int(right.get("multilabel_present_assignments", 0)),
        "labels": dict(sorted(labels.items())),
        "workflow_support_review": workflow,
    }


def write_uncoded_csv(jsonl_path: Path, target: Path) -> tuple[int, set[str]]:
    fields = [
        "canonical_id", "title", "abstract", "doi", "source_url", "coding_source",
        "coding_order", "abstract_status",
    ]
    seen: set[str] = set()
    count = 0
    with target.open("x", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for row in load_jsonl(jsonl_path):
            record_id = row["canonical_id"]
            if record_id in seen:
                raise ValueError(f"duplicate uncoded ID: {record_id}")
            seen.add(record_id)
            writer.writerow({
                "canonical_id": record_id, "title": row.get("title", ""),
                "abstract": row.get("abstract", ""), "doi": row.get("doi") or "",
                "source_url": row.get("source_url") or "", "coding_source": row.get("coding_source", ""),
                "coding_order": row.get("coding_order", ""),
                "abstract_status": "MISSING" if not str(row.get("abstract", "")).strip() else "PRESENT",
            })
            count += 1
    return count, seen


def write_workbook(csv_path: Path, target: Path) -> dict[str, Any]:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    workbook = Workbook()
    instructions = workbook.active
    instructions.title = "Instructions"
    instructions.append(["Combined provisional category coding"])
    instructions.append(["These title-and-abstract assignments are provisional and do not establish final eligibility."])
    instructions.append(["The 646 initial precision candidates and remaining-queue candidates are identified in Candidate origin."])
    instructions.append(["UNKNOWN or UNSUPPORTED coding flags do not rewrite preserved screening outcomes."])
    sheet = workbook.create_sheet("Categorized candidates")
    with csv_path.open(encoding="utf-8", newline="") as handle:
        for row in csv.reader(handle):
            sheet.append(row)
    header = [cell.value for cell in sheet[1]]
    dark = PatternFill("solid", fgColor="263B50")
    input_fill = PatternFill("solid", fgColor="FFF4CC")
    for cell in sheet[1]:
        cell.fill = dark
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for field, formula in (
        ("human_category_review", '"CONFIRMED,REVISED,UNCERTAIN"'),
        ("human_final_eligibility", '"INCLUDE,EXCLUDE,DEFER,NOT_ASSESSED"'),
    ):
        if field not in header:
            raise ValueError(f"missing editable workbook field: {field}")
        column = header.index(field) + 1
        validation = DataValidation(type="list", formula1=formula, allow_blank=True)
        sheet.add_data_validation(validation)
        validation.add(f"{sheet.cell(2, column).coordinate}:{sheet.cell(max(2, sheet.max_row), column).coordinate}")
        for row_number in range(2, sheet.max_row + 1):
            sheet.cell(row_number, column).fill = input_fill
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = {
        "candidate_origin": 30, "coding_stage_run": 18, "canonical_id": 34,
        "title": 55, "abstract": 85, "doi": 28, "source_url": 42,
        "category_evidence_ids": 60, "coding_rationale": 70,
        "human_category_review": 24, "human_final_eligibility": 24, "human_notes": 45,
    }
    for index, field in enumerate(header, 1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = widths.get(field, 24)
    sheet.row_dimensions[1].height = 34
    instructions.column_dimensions["A"].width = 110
    for row in range(1, 5):
        instructions.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="top")
    workbook.save(target)
    verified = load_workbook(target, read_only=False, data_only=False)
    if verified.sheetnames != ["Instructions", "Categorized candidates"]:
        raise ValueError("saved workbook sheet layout drift")
    verified_sheet = verified["Categorized candidates"]
    result = {
        "sheets": verified.sheetnames,
        "candidate_rows": verified_sheet.max_row - 1,
        "columns": verified_sheet.max_column,
        "freeze_panes": str(verified_sheet.freeze_panes),
        "data_validations": len(verified_sheet.data_validations.dataValidation),
    }
    verified.close()
    return result


def build_handoff(
    *, parent_deployment: Path, continuation_run: Path, original_coding_seed: Path, output_dir: Path
) -> dict[str, Any]:
    parent_run = parent_deployment / "results" / "run-v1"
    parent_report = json.loads((parent_run / "run_report.json").read_text(encoding="utf-8"))
    continuation_report = json.loads((continuation_run / "run_report.json").read_text(encoding="utf-8"))
    if not (continuation_run / "terminal_status.json").exists():
        raise ValueError("continuation is not terminal")
    output_dir.mkdir(parents=True, exist_ok=False)

    group_names = ("advance", "deferred", "excluded", "background", "failed", "ambiguous")
    group_counts = {}
    group_ids = {}
    for group in group_names:
        count, ids = combine_csv(
            [parent_run / f"screening_{group}.csv", continuation_run / f"screening_{group}.csv"],
            output_dir / f"combined_{group}.csv",
        )
        group_counts[group] = count
        group_ids[group] = ids
    outcome_union = group_ids["advance"] | group_ids["deferred"] | group_ids["excluded"]
    if len(outcome_union) != sum(group_counts[name] for name in ("advance", "deferred", "excluded")):
        raise ValueError("screening outcome groups overlap")
    if not group_ids["background"].issubset(group_ids["excluded"]):
        raise ValueError("background records are not a subset of exclusions")

    unprocessed_jsonl = continuation_run / "screening_unprocessed.jsonl"
    unprocessed_count = sum(1 for _ in load_jsonl(unprocessed_jsonl))
    unprocessed_ids = {row["canonical_id"] for row in load_jsonl(unprocessed_jsonl)}
    if len(unprocessed_ids) != unprocessed_count:
        raise ValueError("duplicate continuation unprocessed IDs")
    coverage_total = len(outcome_union) + group_counts["failed"] + group_counts["ambiguous"] + unprocessed_count
    if coverage_total != 126168:
        raise ValueError(f"combined screening coverage does not reconcile: {coverage_total}")
    if outcome_union & (group_ids["failed"] | group_ids["ambiguous"] | unprocessed_ids):
        raise ValueError("valid and non-valid screening groups overlap")

    initial_ids = {row["canonical_id"] for row in load_jsonl(original_coding_seed)}
    if len(initial_ids) != 646:
        raise ValueError("initial 646 coding membership drift")
    categorized_count, categorized_ids, categorized_fields = combine_categorized(
        parent_csv=parent_run / "categorized_candidates.csv",
        continuation_csv=continuation_run / "categorized_candidates.csv",
        initial_ids=initial_ids,
        target=output_dir / "combined_categorized_candidates.csv",
    )
    uncoded_count, uncoded_ids = write_uncoded_csv(
        continuation_run / "uncoded_survivors.jsonl",
        output_dir / "combined_uncoded_candidates.csv",
    )
    if categorized_ids & uncoded_ids:
        raise ValueError("categorized and uncoded candidate memberships overlap")
    if categorized_count + uncoded_count != group_counts["advance"] + 646:
        raise ValueError("candidate coding coverage does not reconcile")
    initial_coded = len(categorized_ids & initial_ids)
    initial_uncoded = len(uncoded_ids & initial_ids)
    if initial_coded + initial_uncoded != 646:
        raise ValueError("initial 646 coding coverage does not reconcile")

    warnings = e7_warning_summary([
        parent_run / "screening_records", continuation_run / "screening_records"
    ])
    category_summary = sum_category_summaries(
        parent_report["category_summary"], continuation_report["category_summary"]
    )
    if category_summary["unique_coded_papers"] != categorized_count:
        raise ValueError("combined category unique-paper count drift")
    workbook_path = output_dir / "combined_categorized_candidates.xlsx"
    workbook_check = write_workbook(output_dir / "combined_categorized_candidates.csv", workbook_path)
    if workbook_check["candidate_rows"] != categorized_count:
        raise ValueError("combined workbook row count drift")

    coverage = {
        "original_queue_records": 126168,
        "valid_screenings": len(outcome_union),
        "advance": group_counts["advance"], "defer": group_counts["deferred"],
        "excluded": group_counts["excluded"], "background_subset_of_excluded": group_counts["background"],
        "failed": group_counts["failed"], "ambiguous": group_counts["ambiguous"],
        "still_unprocessed": unprocessed_count, "reconciliation_total": coverage_total,
        "initial_precision_candidates": 646, "initial_coded": initial_coded,
        "initial_uncoded": initial_uncoded,
        "remaining_queue_advance_candidates": group_counts["advance"],
        "all_candidate_records": group_counts["advance"] + 646,
        "categorized_candidates": categorized_count, "uncoded_candidates": uncoded_count,
    }
    spending = {
        "historical_before_remaining_campaign_usd": parent_report["historical_cumulative_conservative_cost_usd"],
        "parent_terminal_cumulative_usd": parent_report["cumulative_conservative_cost_usd"],
        "continuation_terminal_cumulative_usd": continuation_report["cumulative_conservative_cost_usd"],
        "terminal_outstanding_reservations_usd": continuation_report["outstanding_reservations_usd"],
        "hard_cumulative_cap_usd": continuation_report["hard_cumulative_spending_cap_usd"],
        "continuation_screening_usage": continuation_report["screening"]["usage"],
        "continuation_coding_usage": continuation_report["coding"]["usage"],
        "cumulative_usage": continuation_report["cumulative_usage"],
        "stopping_reason": continuation_report["stopping_reason"],
    }
    report = {
        "artifact_class": "title_abstract_remaining_campaign_combined_handoff",
        "status": "TERMINAL_COMBINED_VIEW",
        "coverage": coverage,
        "e7_warnings_by_screening_outcome": warnings,
        "category_summary": category_summary,
        "spending": spending,
        "workbook_verification": workbook_check,
        "candidate_origin_preserved": True,
        "background_double_counted_as_outcome": False,
        "category_assignments_are_provisional": True,
        "final_eligibility_established": False,
        "paths": {
            "categorized_csv": str(output_dir / "combined_categorized_candidates.csv"),
            "categorized_workbook": str(workbook_path),
            "uncoded_csv": str(output_dir / "combined_uncoded_candidates.csv"),
        },
    }
    create_json(output_dir / "combined_report.json", report)
    artifacts = {}
    for path in sorted(output_dir.iterdir()):
        if path.is_file() and path.name != "package_manifest.json":
            artifacts[path.name] = {"sha256": sha256_file(path), "bytes": path.stat().st_size}
    create_json(output_dir / "package_manifest.json", {
        "artifact_class": "title_abstract_remaining_campaign_combined_handoff_package",
        "coverage": coverage, "artifacts": artifacts, "staging_only": True,
    })
    return report


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--parent-deployment", type=Path, required=True)
    parser.add_argument("--continuation-run", type=Path, required=True)
    parser.add_argument("--original-coding-seed", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    report = build_handoff(
        parent_deployment=args.parent_deployment, continuation_run=args.continuation_run,
        original_coding_seed=args.original_coding_seed, output_dir=args.output_dir,
    )
    print(json.dumps(report, indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
